from datetime import datetime, time, timedelta
from io import BytesIO
from zipfile import ZipFile

from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory, TestCase
from django.utils import timezone
from rest_framework.test import APITestCase
from openpyxl import Workbook

from accounts.models import User
from audit.models import ActivityLog, DataSubmission, RecycleBin, SubmissionNotification
from audit.services import log_action, purge_recycle_entry, restore_recycle_entry, soft_delete_to_recycle
from citosis_pro.common import RecycleItemTypeChoices, RoleChoices, UserStatusChoices
from tourism.models import TourismRecord, Visitor


def create_user(email, username, role=RoleChoices.USER, status=UserStatusChoices.ACTIVE):
    return User.objects.create_user(
        email=email,
        password='pass1234',
        username=username,
        name='Test User',
        role=role,
        status=status,
        email_verified_at=timezone.now(),
    )


def build_excel_upload(filename, sheets):
    workbook = Workbook()
    worksheet = workbook.active
    first_sheet = True

    for title, rows in sheets:
        if first_sheet:
            worksheet.title = title
            target_sheet = worksheet
            first_sheet = False
        else:
            target_sheet = workbook.create_sheet(title=title)
        for row in rows:
            target_sheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        filename,
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def build_csv_upload(filename, content):
    return SimpleUploadedFile(
        filename,
        content.encode('utf-8'),
        content_type='text/csv',
    )


def set_model_timestamps(instance, created_at, updated_at=None):
    instance.__class__._default_manager.filter(pk=instance.pk).update(
        created_at=created_at,
        updated_at=updated_at or created_at,
    )
    instance.refresh_from_db()
    return instance


class AuditServiceTests(TestCase):
    def setUp(self):
        self.actor = create_user('admin@example.com', 'admin', role=RoleChoices.SUPER_ADMIN)

    def test_log_action_captures_request_meta(self):
        request = RequestFactory().get('/', HTTP_X_FORWARDED_FOR='10.0.0.1', HTTP_USER_AGENT='AuditTestAgent')
        log_action(self.actor, 'Test action', 'Did a thing.', request)

        log_entry = ActivityLog.objects.get()
        self.assertEqual(log_entry.user, self.actor)
        self.assertEqual(log_entry.action, 'Test action')
        self.assertEqual(log_entry.details, 'Did a thing.')
        self.assertEqual(log_entry.ip_address, '10.0.0.1')
        self.assertEqual(log_entry.user_agent, 'AuditTestAgent')

    def test_soft_delete_to_recycle_marks_user_inactive_and_creates_entry(self):
        user = create_user('user@example.com', 'user')
        soft_delete_to_recycle(user, RecycleItemTypeChoices.USER, self.actor)

        user.refresh_from_db()
        self.assertIsNotNone(user.deleted_at)
        self.assertEqual(user.status, UserStatusChoices.INACTIVE)
        self.assertFalse(user.is_active)
        self.assertNotEqual(user.email, 'user@example.com')
        self.assertTrue(user.email.endswith('@deleted.citosis.local'))

        entry = RecycleBin.objects.get(item_id=str(user.pk))
        self.assertEqual(entry.item_type, RecycleItemTypeChoices.USER)
        self.assertEqual(entry.deleted_by, self.actor)
        self.assertEqual(entry.item_data['email'], 'user@example.com')

    def test_restore_recycle_entry_reactivates_user(self):
        user = create_user('restore@example.com', 'restore')
        soft_delete_to_recycle(user, RecycleItemTypeChoices.USER, self.actor)
        entry = RecycleBin.objects.get(item_id=str(user.pk))

        restore_recycle_entry(entry, self.actor)

        user = User.all_objects.get(pk=user.pk)
        self.assertIsNone(user.deleted_at)
        self.assertEqual(user.status, UserStatusChoices.ACTIVE)
        self.assertTrue(user.is_active)

        entry.refresh_from_db()
        self.assertIsNotNone(entry.restored_at)
        self.assertEqual(entry.restored_by, self.actor)

    def test_purge_recycle_entry_deletes_soft_deleted_instance(self):
        user = create_user('purge@example.com', 'purge')
        soft_delete_to_recycle(user, RecycleItemTypeChoices.USER, self.actor)
        entry = RecycleBin.objects.get(item_id=str(user.pk))

        purge_recycle_entry(entry, self.actor)

        self.assertFalse(RecycleBin.objects.filter(pk=entry.pk).exists())
        self.assertFalse(User.all_objects.filter(pk=user.pk).exists())

    def test_purge_recycle_entry_deletes_user_referenced_by_tourism_records(self):
        user = create_user('owner@example.com', 'owner')
        TourismRecord.objects.create(
            name='Capistrano',
            location='Malaybalay',
            category='Nature',
            created_by=user,
        )
        Visitor.objects.create(
            name='Visitor One',
            place='Capistrano',
            origin='Malaybalay',
            visit_date=timezone.now(),
            created_by=user,
        )
        soft_delete_to_recycle(user, RecycleItemTypeChoices.USER, self.actor)
        entry = RecycleBin.objects.get(item_id=str(user.pk))

        purge_recycle_entry(entry, self.actor)

        self.assertFalse(RecycleBin.objects.filter(pk=entry.pk).exists())
        self.assertFalse(User.all_objects.filter(pk=user.pk).exists())
        self.assertIsNone(TourismRecord.objects.get(name='Capistrano').created_by)
        self.assertIsNone(Visitor.objects.get(name='Visitor One').created_by)


class AuditApiTests(APITestCase):
    def setUp(self):
        self.admin = create_user('apiadmin@example.com', 'apiadmin', role=RoleChoices.SUPER_ADMIN)
        self.reviewer = create_user('reviewer@example.com', 'reviewer', role=RoleChoices.REVIEWER)
        self.viewer = create_user('viewer@example.com', 'viewer', role=RoleChoices.VIEWER)
        self.staff = create_user('apistaff@example.com', 'apistaff', role=RoleChoices.USER)

    def test_dashboard_overview_returns_payload(self):
        self.client.force_authenticate(user=self.admin)
        response = self.client.get('/api/dashboard/overview/')

        self.assertEqual(response.status_code, 200)
        self.assertIn('stats', response.data)
        self.assertIn('intelligence', response.data)
        self.assertIn('destination_popularity', response.data)
        self.assertIn('recent_activity', response.data)
        self.assertIn('recent_records', response.data)
        self.assertIn('summary', response.data)
        self.assertIn('records', response.data['stats'])

    def test_dashboard_overview_returns_admin_intelligence_metrics(self):
        now = timezone.make_aware(
            datetime.combine(timezone.localdate(), time(hour=12)),
            timezone.get_current_timezone(),
        )
        record = TourismRecord.objects.create(
            name='River Park',
            location='Malaybalay City',
            category='Park',
            created_by=self.admin,
        )

        visitor_today_one = Visitor.objects.create(
            name='Ana Cruz',
            place='River Park',
            origin='Malaybalay City',
            visit_date=now,
            tourism_record=record,
            created_by=self.admin,
        )
        visitor_today_two = Visitor.objects.create(
            name='Ben Lopez',
            place='River Park',
            origin='Valencia City',
            visit_date=now,
            tourism_record=record,
            created_by=self.admin,
        )
        visitor_older = Visitor.objects.create(
            name='Cara Lim',
            place='Bean Zone',
            origin='Cagayan de Oro',
            visit_date=now - timedelta(days=2),
            created_by=self.admin,
        )
        set_model_timestamps(visitor_today_one, now - timedelta(hours=3))
        set_model_timestamps(visitor_today_two, now - timedelta(hours=2))
        set_model_timestamps(visitor_older, now - timedelta(days=2))

        pending_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('pending.xlsx', b'pending', content_type='application/vnd.ms-excel'),
            original_filename='pending.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_PENDING,
        )
        in_review_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('review.xlsx', b'review', content_type='application/vnd.ms-excel'),
            original_filename='review.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_IN_REVIEW,
        )
        approved_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('approved.xlsx', b'approved', content_type='application/vnd.ms-excel'),
            original_filename='approved.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_APPROVED,
        )
        rejected_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('rejected.xlsx', b'rejected', content_type='application/vnd.ms-excel'),
            original_filename='rejected.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_REJECTED,
        )
        draft_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('draft.xlsx', b'draft', content_type='application/vnd.ms-excel'),
            original_filename='draft.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_DRAFT,
        )

        set_model_timestamps(pending_submission, now - timedelta(hours=1))
        set_model_timestamps(in_review_submission, now - timedelta(days=1))
        set_model_timestamps(
            approved_submission,
            now - timedelta(days=2, hours=4),
            now - timedelta(days=2, hours=3),
        )
        set_model_timestamps(
            rejected_submission,
            now - timedelta(days=3, hours=6),
            now - timedelta(days=3, hours=3),
        )
        set_model_timestamps(draft_submission, now - timedelta(minutes=30))

        self.client.force_authenticate(user=self.admin)
        response = self.client.get('/api/dashboard/overview/')

        self.assertEqual(response.status_code, 200)
        metrics = {item['key']: item for item in response.data['intelligence']['metrics']}
        self.assertEqual(metrics['submissions_today']['value'], 1)
        self.assertEqual(metrics['pending_reviews']['value'], 2)
        self.assertEqual(metrics['rejection_rate']['value'], 50.0)
        self.assertEqual(metrics['avg_review_time']['value'], 7200)
        self.assertEqual(metrics['avg_review_time']['display_value'], '2h')

        charts = response.data['intelligence']['charts']
        self.assertEqual(charts['submissions']['total'], 4)
        self.assertEqual(len(charts['submissions']['points']), 7)
        self.assertEqual(charts['visitors']['delta'], 3)
        self.assertTrue(charts['destinations']['items'])
        self.assertEqual(charts['destinations']['items'][0]['label'], 'River Park')

    def test_dashboard_overview_requires_admin(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/dashboard/overview/')

        self.assertEqual(response.status_code, 403)

    def test_viewer_can_access_dashboard_overview(self):
        self.client.force_authenticate(user=self.viewer)

        response = self.client.get('/api/dashboard/overview/')

        self.assertEqual(response.status_code, 200)

    def test_recycle_bin_list_requires_admin(self):
        self.client.force_authenticate(user=self.staff)
        response = self.client.get('/api/recycle-bin/')

        self.assertEqual(response.status_code, 403)

    def test_user_can_preview_excel_submission_with_sheet_selection_and_validation_feedback(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_excel_upload(
            'preview.xlsx',
            [
                (
                    'Overview',
                    [
                        ['Place', 'Location', 'Category'],
                        ['Bean Zone', 'Capitol Grounds', 'Event'],
                    ],
                ),
                (
                    'Visitors',
                    [
                        ['Date', 'Name', 'Place'],
                        ['not-a-date', 'Ana', 'River Park'],
                        ['not-a-date', 'Ana', 'River Park'],
                    ],
                ),
            ],
        )

        response = self.client.post(
            '/api/data-submissions/preview/',
            {'file': upload, 'sheet_name': 'Visitors'},
            format='multipart',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['sheet_name'], 'Visitors')
        self.assertEqual(response.data['selected_sheet_name'], 'Visitors')
        self.assertEqual(len(response.data['sheets']), 2)
        self.assertEqual(response.data['validation']['missing_required_columns'], ['Origin'])
        self.assertEqual(response.data['validation']['duplicate_count'], 1)
        self.assertIn('Date', response.data['validation']['invalid_format_messages'][0])

    def test_user_can_preview_excel_submission_with_manual_column_mapping(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_excel_upload(
            'mapped-preview.xlsx',
            [
                (
                    'Visitors',
                    [
                        ['Visit Day', 'Guest', 'Attraction', 'Hometown'],
                        ['2026-04-02', 'Ana', 'River Park', 'Malaybalay City'],
                    ],
                ),
            ],
        )

        response = self.client.post(
            '/api/data-submissions/preview/',
            {
                'file': upload,
                'sheet_name': 'Visitors',
                'mapping_profile_key': 'visitor',
                'column_mapping': '{"Date":"Visit Day","Name":"Guest","Place":"Attraction","Origin":"Hometown"}',
            },
            format='multipart',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['mapping_profile_key'], 'visitor')
        self.assertEqual(response.data['validation']['missing_required_columns'], [])
        self.assertEqual(response.data['column_mapping']['Name'], 'Guest')
        self.assertEqual(response.data['validation']['applied_column_mapping']['Place'], 'Attraction')

    def test_user_can_preview_csv_submission_with_validation_feedback(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_csv_upload(
            'visitor-preview.csv',
            '\n'.join([
                'Date,Name,Place,Origin',
                'not-a-date,Ana,River Park,Malaybalay City',
                'not-a-date,Ana,River Park,Malaybalay City',
            ]),
        )

        response = self.client.post(
            '/api/data-submissions/preview/',
            {'file': upload},
            format='multipart',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['sheet_name'], 'visitor-preview')
        self.assertEqual(response.data['selected_sheet_name'], 'visitor-preview')
        self.assertEqual(len(response.data['sheets']), 1)
        self.assertEqual(response.data['headers'], ['Date', 'Name', 'Place', 'Origin'])
        self.assertEqual(response.data['validation']['profile_key'], 'visitor')
        self.assertEqual(response.data['validation']['duplicate_count'], 1)
        self.assertIn('Date', response.data['validation']['invalid_format_messages'][0])

    def test_user_can_send_excel_data_submission(self):
        self.client.force_authenticate(user=self.staff)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Submission'
        worksheet.append(['Name', 'Count'])
        worksheet.append(['River Park', 12])
        buffer = BytesIO()
        workbook.save(buffer)
        upload = SimpleUploadedFile(
            'submission.xlsx',
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        response = self.client.post('/api/data-submissions/', {'file': upload}, format='multipart')

        admin_alert = SubmissionNotification.objects.get(
            submission__submitted_by=self.staff,
            submission__original_filename='submission.xlsx',
            recipient=self.admin,
        )
        self.assertEqual(response.status_code, 201)
        self.assertTrue(DataSubmission.objects.filter(submitted_by=self.staff, original_filename='submission.xlsx').exists())
        self.assertEqual(response.data['sheet_name'], 'Submission')
        self.assertEqual(response.data['headers'], ['Name', 'Count'])
        self.assertEqual(response.data['review_history_count'], 1)
        self.assertEqual(admin_alert.title, 'New submission received')
        self.assertTrue(any(self.admin.email in email.to for email in mail.outbox))
        self.assertTrue(any(self.reviewer.email in email.to for email in mail.outbox))
        self.assertTrue(any('CITOSIS PRO admin notification' in email.subject for email in mail.outbox))

    def test_user_can_send_csv_data_submission(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_csv_upload(
            'submission.csv',
            '\n'.join([
                'Name,Count',
                'River Park,12',
            ]),
        )

        response = self.client.post('/api/data-submissions/', {'file': upload}, format='multipart')

        admin_alert = SubmissionNotification.objects.get(
            submission__submitted_by=self.staff,
            submission__original_filename='submission.csv',
            recipient=self.admin,
        )
        self.assertEqual(response.status_code, 201)
        self.assertTrue(DataSubmission.objects.filter(submitted_by=self.staff, original_filename='submission.csv').exists())
        self.assertEqual(response.data['sheet_name'], 'submission')
        self.assertEqual(response.data['headers'], ['Name', 'Count'])
        self.assertEqual(response.data['preview_rows'], [['River Park', '12']])
        self.assertEqual(response.data['review_history_count'], 1)
        self.assertEqual(admin_alert.title, 'New submission received')

    def test_high_priority_submission_creates_priority_admin_alert(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_excel_upload(
            'visitor-batch.xlsx',
            [
                (
                    'Visitors',
                    [
                        ['Date', 'Name', 'Place', 'Origin'],
                        ['2026-04-02', 'Ana', 'River Park', 'Malaybalay City'],
                    ],
                ),
            ],
        )

        response = self.client.post(
            '/api/data-submissions/',
            {
                'file': upload,
                'sheet_name': 'Visitors',
                'notes': 'Urgent file for immediate review.',
            },
            format='multipart',
        )

        alert = SubmissionNotification.objects.get(
            submission__submitted_by=self.staff,
            submission__original_filename='visitor-batch.xlsx',
            recipient=self.admin,
        )
        self.assertEqual(response.status_code, 201)
        self.assertTrue(response.data['is_high_priority'])
        self.assertEqual(alert.title, 'High-priority file uploaded')
        self.assertIn('submission notes', alert.message)
        self.assertTrue(any('High-priority file uploaded' in email.subject for email in mail.outbox))

    def test_user_can_send_excel_data_submission_with_mapping_metadata(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_excel_upload(
            'mapped-submission.xlsx',
            [
                (
                    'Visitors',
                    [
                        ['Visit Day', 'Guest', 'Attraction', 'Hometown'],
                        ['2026-04-02', 'Ana', 'River Park', 'Malaybalay City'],
                    ],
                ),
            ],
        )

        response = self.client.post(
            '/api/data-submissions/',
            {
                'file': upload,
                'sheet_name': 'Visitors',
                'mapping_profile_key': 'visitor',
                'column_mapping': '{"Date":"Visit Day","Name":"Guest","Place":"Attraction","Origin":"Hometown"}',
            },
            format='multipart',
        )

        submission = DataSubmission.objects.get(original_filename='mapped-submission.xlsx')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['mapping_profile_key'], 'visitor')
        self.assertEqual(response.data['mapping_profile_label'], 'Visitor Sheet')
        self.assertEqual(submission.column_mapping['Origin'], 'Hometown')

    def test_user_can_download_submission_template(self):
        self.client.force_authenticate(user=self.staff)

        response = self.client.get('/api/data-submission-templates/visitor/')

        self.assertEqual(response.status_code, 200)
        self.assertIn('visitor-upload-template.xlsx', response['Content-Disposition'])
        self.assertGreater(len(response.content), 0)

    def test_user_can_send_a_specific_sheet_from_excel_submission(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_excel_upload(
            'submission-multi.xlsx',
            [
                (
                    'SheetOne',
                    [
                        ['Name', 'Count'],
                        ['River Park', 12],
                    ],
                ),
                (
                    'ChosenSheet',
                    [
                        ['Place', 'Location', 'Category'],
                        ['Bean Zone', 'Capitol Grounds', 'Event'],
                    ],
                ),
            ],
        )

        response = self.client.post(
            '/api/data-submissions/',
            {'file': upload, 'sheet_name': 'ChosenSheet'},
            format='multipart',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['sheet_name'], 'ChosenSheet')
        self.assertEqual(response.data['headers'], ['Place', 'Location', 'Category'])
        self.assertTrue(
            DataSubmission.objects.filter(
                submitted_by=self.staff,
                original_filename='submission-multi.xlsx',
                sheet_name='ChosenSheet',
            ).exists()
        )

    def test_user_can_save_excel_submission_as_draft_with_notes(self):
        self.client.force_authenticate(user=self.staff)
        upload = build_excel_upload(
            'draft-report.xlsx',
            [
                (
                    'April',
                    [
                        ['Place', 'Location', 'Category'],
                        ['Bean Zone', 'Capitol Grounds', 'Event'],
                    ],
                ),
            ],
        )

        response = self.client.post(
            '/api/data-submissions/',
            {
                'file': upload,
                'sheet_name': 'April',
                'save_mode': 'draft',
                'notes': 'Data is incomplete for April.',
            },
            format='multipart',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['status'], DataSubmission.STATUS_DRAFT)
        self.assertEqual(response.data['submission_notes'], 'Data is incomplete for April.')
        self.assertEqual(response.data['version_number'], 1)

    def test_user_can_send_saved_draft_later(self):
        draft = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('draft.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='draft.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            submission_notes='Please prioritize this.',
            status=DataSubmission.STATUS_DRAFT,
        )
        self.client.force_authenticate(user=self.staff)

        response = self.client.post(
            f'/api/data-submissions/{draft.pk}/send/',
            {'notes': 'Please prioritize this.'},
            format='json',
        )

        draft.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['status'], DataSubmission.STATUS_PENDING)
        self.assertEqual(draft.status, DataSubmission.STATUS_PENDING)

    def test_version_number_increments_for_same_file_and_sheet(self):
        self.client.force_authenticate(user=self.staff)
        first_upload = build_excel_upload(
            'monthly-report.xlsx',
            [
                (
                    'Visitors',
                    [
                        ['Date', 'Name', 'Place', 'Origin'],
                        ['2026-04-01', 'Ana', 'River Park', 'Malaybalay'],
                    ],
                ),
            ],
        )
        second_upload = build_excel_upload(
            'monthly-report.xlsx',
            [
                (
                    'Visitors',
                    [
                        ['Date', 'Name', 'Place', 'Origin'],
                        ['2026-04-02', 'Ben', 'River Park', 'Valencia'],
                    ],
                ),
            ],
        )

        first_response = self.client.post(
            '/api/data-submissions/',
            {'file': first_upload, 'sheet_name': 'Visitors'},
            format='multipart',
        )
        second_response = self.client.post(
            '/api/data-submissions/',
            {'file': second_upload, 'sheet_name': 'Visitors'},
            format='multipart',
        )

        self.assertEqual(first_response.status_code, 201)
        self.assertEqual(second_response.status_code, 201)
        self.assertEqual(first_response.data['version_number'], 1)
        self.assertEqual(second_response.data['version_number'], 2)
        self.assertEqual(second_response.data['version_label'], 'v2')

    def test_admin_can_view_all_data_submissions(self):
        DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('existing.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='existing.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.get('/api/data-submissions/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]['submitted_by'], self.staff.id)

    def test_admin_list_excludes_draft_submissions(self):
        DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('draft-only.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='draft-only.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_DRAFT,
        )
        DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('pending.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='pending.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_PENDING,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.get('/api/data-submissions/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]['original_filename'], 'pending.xlsx')

    def test_admin_can_get_submission_workspace_with_issue_details(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('workspace.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='workspace.xlsx',
            sheet_name='Visitors',
            headers=['Date', 'Name', 'Place', 'Origin'],
            preview_rows=[
                ['bad-date', 'Ana', 'River Park', ''],
                ['bad-date', 'Ana', 'River Park', ''],
            ],
            total_rows=2,
            total_columns=4,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.get(f'/api/data-submissions/{submission.pk}/workspace/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['workspace_preview']['source_submission_id'], submission.pk)
        self.assertIn('Origin', response.data['workspace_preview']['validation']['missing_value_messages'][0])
        self.assertEqual(response.data['workspace_preview']['validation']['duplicate_row_indexes'], [0, 1])
        self.assertTrue(
            any(
                issue['column_index'] == 0 and issue['tone'] == 'error'
                for issue in response.data['workspace_preview']['validation']['cell_issues']
            )
        )

    def test_admin_can_update_submission_workspace_and_mapping(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('workspace-edit.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='workspace-edit.xlsx',
            sheet_name='Visitors',
            headers=['Visit Day', 'Guest', 'Attraction', 'Hometown'],
            preview_rows=[['bad-date', 'Ana', 'River Park', 'Malaybalay City']],
            total_rows=1,
            total_columns=4,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.patch(
            f'/api/data-submissions/{submission.pk}/workspace/',
            {
                'headers': ['Visit Day', 'Guest', 'Attraction', 'Hometown'],
                'preview_rows': [['2026-04-02', 'Ana Cruz', 'River Park', 'Malaybalay City']],
                'mapping_profile_key': 'visitor',
                'column_mapping': {
                    'Date': 'Visit Day',
                    'Name': 'Guest',
                    'Place': 'Attraction',
                    'Origin': 'Hometown',
                },
            },
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(submission.mapping_profile_key, 'visitor')
        self.assertEqual(submission.column_mapping['Name'], 'Guest')
        self.assertEqual(submission.preview_rows[0][0], '2026-04-02')
        self.assertEqual(submission.last_reviewed_by, self.admin)
        self.assertEqual(submission.review_history[-1]['action_key'], 'workspace_saved')
        self.assertEqual(response.data['workspace_preview']['validation']['missing_required_columns'], [])
        self.assertEqual(response.data['workspace_preview']['review_history'][0]['action_key'], 'workspace_saved')
        self.assertEqual(response.data['submission']['mapping_profile_label'], 'Visitor Sheet')

    def test_admin_workspace_response_includes_version_comparison_for_resubmission(self):
        DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('resubmission.xlsx', b'v1', content_type='application/vnd.ms-excel'),
            original_filename='resubmission.xlsx',
            sheet_name='Visitors',
            headers=['Date', 'Name', 'Place', 'Origin'],
            preview_rows=[['2026-04-01', 'Ana', 'River Park', 'Malaybalay City']],
            total_rows=1,
            total_columns=4,
            version_number=1,
        )
        latest_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('resubmission.xlsx', b'v2', content_type='application/vnd.ms-excel'),
            original_filename='resubmission.xlsx',
            sheet_name='Visitors',
            headers=['Visit Date', 'Name', 'Place', 'Origin'],
            preview_rows=[['2026-04-02', 'Ana', 'River Park', 'Valencia City']],
            total_rows=1,
            total_columns=4,
            version_number=2,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.get(f'/api/data-submissions/{latest_submission.pk}/workspace/')

        comparison = response.data['workspace_preview']['version_comparison']
        self.assertEqual(response.status_code, 200)
        self.assertTrue(comparison['has_previous_version'])
        self.assertEqual(comparison['previous_version_label'], 'v1')
        self.assertEqual(comparison['current_version_label'], 'v2')
        self.assertEqual(comparison['metrics']['header_change_count'], 1)
        self.assertEqual(comparison['metrics']['cell_changed_count'], 2)
        self.assertTrue(comparison['items'])

    def test_admin_can_validate_submission_workspace_without_saving(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('workspace-validate.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='workspace-validate.xlsx',
            sheet_name='Visitors',
            headers=['Date', 'Name', 'Place', 'Origin'],
            preview_rows=[['2026-04-01', 'Ana', 'River Park', 'Malaybalay City']],
            total_rows=1,
            total_columns=4,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.post(
            f'/api/data-submissions/{submission.pk}/workspace/validate/',
            {
                'headers': ['Date', 'Name', 'Place', 'Origin'],
                'preview_rows': [['', 'Ana', 'River Park', 'Malaybalay City']],
                'mapping_profile_key': 'visitor',
                'column_mapping': {
                    'Date': 'Date',
                    'Name': 'Name',
                    'Place': 'Place',
                    'Origin': 'Origin',
                },
            },
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertIn('Date', response.data['workspace_preview']['validation']['missing_value_messages'][0])
        self.assertEqual(submission.preview_rows[0][0], '2026-04-01')

    def test_admin_can_move_data_submission_to_in_review(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('review.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='review.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.patch(
            f'/api/data-submissions/{submission.pk}/status/',
            {'status': DataSubmission.STATUS_IN_REVIEW},
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['status'], DataSubmission.STATUS_IN_REVIEW)
        self.assertEqual(submission.status, DataSubmission.STATUS_IN_REVIEW)
        self.assertEqual(submission.last_reviewed_by, self.admin)
        self.assertEqual(submission.review_history[-1]['action_key'], 'status_updated')

    def test_reviewer_can_move_data_submission_to_in_review(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('reviewer-review.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='reviewer-review.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.reviewer)

        response = self.client.patch(
            f'/api/data-submissions/{submission.pk}/status/',
            {'status': DataSubmission.STATUS_IN_REVIEW},
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(submission.status, DataSubmission.STATUS_IN_REVIEW)
        self.assertEqual(submission.last_reviewed_by, self.reviewer)

    def test_viewer_cannot_update_data_submission_status(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('viewer-review.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='viewer-review.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.viewer)

        response = self.client.patch(
            f'/api/data-submissions/{submission.pk}/status/',
            {'status': DataSubmission.STATUS_IN_REVIEW},
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 403)
        self.assertEqual(submission.status, DataSubmission.STATUS_PENDING)

    def test_admin_can_reject_submission_with_feedback_and_notify_user(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('reject.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='reject.xlsx',
            sheet_name='Sheet1',
            headers=['Visitor ID'],
            preview_rows=[['']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.patch(
            f'/api/data-submissions/{submission.pk}/status/',
            {
                'status': DataSubmission.STATUS_REJECTED,
                'admin_feedback': 'Missing visitor ID column.',
            },
            format='json',
        )

        submission.refresh_from_db()
        notification = SubmissionNotification.objects.get(submission=submission, recipient=self.staff)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(submission.status, DataSubmission.STATUS_REJECTED)
        self.assertEqual(submission.admin_feedback, 'Missing visitor ID column.')
        self.assertEqual(notification.status_snapshot, DataSubmission.STATUS_REJECTED)
        self.assertIn('Missing visitor ID column.', notification.message)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn('reject.xlsx', mail.outbox[0].subject)

    def test_admin_can_bulk_approve_submissions(self):
        first_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('bulk-one.xlsx', b'one', content_type='application/vnd.ms-excel'),
            original_filename='bulk-one.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        second_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('bulk-two.xlsx', b'two', content_type='application/vnd.ms-excel'),
            original_filename='bulk-two.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_IN_REVIEW,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.post(
            '/api/data-submissions/bulk-status/',
            {
                'submission_ids': [first_submission.pk, second_submission.pk],
                'status': DataSubmission.STATUS_APPROVED,
            },
            format='json',
        )

        first_submission.refresh_from_db()
        second_submission.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['updated_count'], 2)
        self.assertEqual(first_submission.status, DataSubmission.STATUS_APPROVED)
        self.assertEqual(second_submission.status, DataSubmission.STATUS_APPROVED)
        self.assertEqual(SubmissionNotification.objects.filter(status_snapshot=DataSubmission.STATUS_APPROVED).count(), 2)

    def test_admin_bulk_reject_requires_feedback(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('bulk-reject.xlsx', b'one', content_type='application/vnd.ms-excel'),
            original_filename='bulk-reject.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.post(
            '/api/data-submissions/bulk-status/',
            {
                'submission_ids': [submission.pk],
                'status': DataSubmission.STATUS_REJECTED,
            },
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 400)
        self.assertEqual(submission.status, DataSubmission.STATUS_PENDING)

    def test_admin_can_bulk_export_selected_submissions(self):
        first_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('export-one.xlsx', b'file-one', content_type='application/vnd.ms-excel'),
            original_filename='export-one.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        second_submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('export-two.xlsx', b'file-two', content_type='application/vnd.ms-excel'),
            original_filename='export-two.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.get(
            f'/api/data-submissions/export/?ids={first_submission.pk},{second_submission.pk}',
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn('application/zip', response['Content-Type'])
        with ZipFile(BytesIO(response.content)) as archive:
            names = archive.namelist()
            self.assertIn('manifest.json', names)
            self.assertTrue(any(name.endswith('export-one.xlsx') for name in names))
            self.assertTrue(any(name.endswith('export-two.xlsx') for name in names))

    def test_admin_needs_feedback_for_rejected_submission(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('reject-no-feedback.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='reject-no-feedback.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.patch(
            f'/api/data-submissions/{submission.pk}/status/',
            {'status': DataSubmission.STATUS_REJECTED},
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 400)
        self.assertEqual(submission.status, DataSubmission.STATUS_PENDING)
        self.assertFalse(SubmissionNotification.objects.filter(submission=submission).exists())

    def test_non_admin_cannot_update_data_submission_status(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('locked.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='locked.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        self.client.force_authenticate(user=self.staff)

        response = self.client.patch(
            f'/api/data-submissions/{submission.pk}/status/',
            {'status': DataSubmission.STATUS_APPROVED},
            format='json',
        )

        submission.refresh_from_db()
        self.assertEqual(response.status_code, 403)
        self.assertEqual(submission.status, DataSubmission.STATUS_PENDING)

    def test_reviewer_can_edit_tourism_record_but_cannot_delete_it(self):
        record = TourismRecord.objects.create(
            name='River Park',
            location='Malaybalay City',
            category='Park',
            created_by=self.admin,
        )
        self.client.force_authenticate(user=self.reviewer)

        update_response = self.client.patch(
            f'/api/records/{record.pk}/',
            {'name': 'River Park Updated'},
            format='json',
        )
        delete_response = self.client.delete(f'/api/records/{record.pk}/')

        record.refresh_from_db()
        self.assertEqual(update_response.status_code, 200)
        self.assertEqual(record.name, 'River Park Updated')
        self.assertEqual(delete_response.status_code, 403)

    def test_viewer_cannot_edit_tourism_record(self):
        record = TourismRecord.objects.create(
            name='Bean Zone',
            location='Malaybalay City',
            category='Event',
            created_by=self.admin,
        )
        self.client.force_authenticate(user=self.viewer)

        response = self.client.patch(
            f'/api/records/{record.pk}/',
            {'name': 'Bean Zone Updated'},
            format='json',
        )

        record.refresh_from_db()
        self.assertEqual(response.status_code, 403)
        self.assertEqual(record.name, 'Bean Zone')

    def test_user_can_view_and_mark_submission_notifications_read(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('notify.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='notify.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
            status=DataSubmission.STATUS_NEEDS_REVISION,
            admin_feedback='Please fix the Date values.',
        )
        notification = SubmissionNotification.objects.create(
            recipient=self.staff,
            submission=submission,
            title='Action required on your Excel submission',
            message='Please fix the Date values.',
            status_snapshot=DataSubmission.STATUS_NEEDS_REVISION,
        )
        self.client.force_authenticate(user=self.staff)

        list_response = self.client.get('/api/submission-notifications/')
        read_response = self.client.post(f'/api/submission-notifications/{notification.pk}/read/')

        notification.refresh_from_db()
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(len(list_response.data), 1)
        self.assertEqual(read_response.status_code, 200)
        self.assertTrue(notification.is_read)

    def test_user_can_mark_all_submission_notifications_read(self):
        submission = DataSubmission.objects.create(
            submitted_by=self.staff,
            file=SimpleUploadedFile('notify-all.xlsx', b'placeholder', content_type='application/vnd.ms-excel'),
            original_filename='notify-all.xlsx',
            sheet_name='Sheet1',
            headers=['Column 1'],
            preview_rows=[['Value']],
            total_rows=1,
            total_columns=1,
        )
        SubmissionNotification.objects.create(
            recipient=self.staff,
            submission=submission,
            title='Your Excel submission is now in review',
            message='The admin is reviewing your file.',
            status_snapshot=DataSubmission.STATUS_IN_REVIEW,
        )
        SubmissionNotification.objects.create(
            recipient=self.staff,
            submission=submission,
            title='Action required on your Excel submission',
            message='Please add the visitor ID column.',
            status_snapshot=DataSubmission.STATUS_NEEDS_REVISION,
        )
        self.client.force_authenticate(user=self.staff)

        response = self.client.post('/api/submission-notifications/read-all/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['updated_count'], 2)
        self.assertFalse(SubmissionNotification.objects.filter(recipient=self.staff, is_read=False).exists())

    def test_admin_can_list_data_request_targets(self):
        inactive_user = create_user(
            'inactive-target@example.com',
            'inactive-target',
            status=UserStatusChoices.INACTIVE,
        )
        new_establishment_user = User.objects.create_user(
            email='fresh-establishment@example.com',
            password='pass1234',
            username='fresh-establishment',
            name='Fresh Establishment User',
            office='Fresh Active Establishment',
            status=UserStatusChoices.ACTIVE,
        )
        self.client.force_authenticate(user=self.admin)

        response = self.client.get('/api/data-request-targets/')

        target_ids = {item['id'] for item in response.data}
        target_establishments = {item['office'] for item in response.data}
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.staff.pk, target_ids)
        self.assertIn(new_establishment_user.pk, target_ids)
        self.assertIn('Fresh Active Establishment', target_establishments)
        self.assertNotIn(self.admin.pk, target_ids)
        self.assertNotIn(self.reviewer.pk, target_ids)
        self.assertNotIn(self.viewer.pk, target_ids)
        self.assertNotIn(inactive_user.pk, target_ids)

    def test_admin_can_send_data_request_notification_to_user(self):
        due_date = timezone.localdate() + timedelta(days=7)
        self.client.force_authenticate(user=self.admin)

        response = self.client.post(
            '/api/data-requests/',
            {
                'recipient_ids': [self.staff.pk],
                'title': 'Monthly visitor data needed',
                'message': 'Please upload the April visitor arrivals workbook.',
                'due_date': due_date.isoformat(),
            },
            format='json',
        )

        notification = SubmissionNotification.objects.get(recipient=self.staff)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['created_count'], 1)
        self.assertIsNone(notification.submission_id)
        self.assertEqual(notification.notification_type, SubmissionNotification.TYPE_DATA_REQUEST)
        self.assertEqual(notification.title, 'Monthly visitor data needed')
        self.assertEqual(notification.metadata['due_date'], due_date.isoformat())
        self.assertEqual(notification.created_by, self.admin)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn('Monthly visitor data needed', mail.outbox[0].subject)

        self.client.force_authenticate(user=self.staff)
        list_response = self.client.get('/api/submission-notifications/')

        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(list_response.data[0]['notification_type'], SubmissionNotification.TYPE_DATA_REQUEST)
        self.assertIsNone(list_response.data[0]['submission_id'])
        self.assertEqual(list_response.data[0]['created_by_name'], self.admin.name)

    def test_admin_can_view_data_request_history(self):
        due_date = timezone.localdate() + timedelta(days=3)
        self.client.force_authenticate(user=self.admin)
        self.client.post(
            '/api/data-requests/',
            {
                'recipient_ids': [self.staff.pk],
                'title': 'Weekly occupancy needed',
                'message': 'Please upload this week\'s occupancy file.',
                'due_date': due_date.isoformat(),
            },
            format='json',
        )

        response = self.client.get('/api/data-requests/history/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]['notification_type'], SubmissionNotification.TYPE_DATA_REQUEST)
        self.assertEqual(response.data[0]['title'], 'Weekly occupancy needed')
        self.assertEqual(response.data[0]['recipient'], self.staff.pk)
        self.assertEqual(response.data[0]['recipient_email'], self.staff.email)
        self.assertEqual(response.data[0]['recipient_office'], self.staff.office)
        self.assertEqual(response.data[0]['metadata']['due_date'], due_date.isoformat())

    def test_admin_can_send_data_request_notification_to_establishment(self):
        first_staff = create_user('farm-one@example.com', 'farm-one')
        first_staff.office = 'New Active Farm'
        first_staff.save()
        second_staff = create_user('farm-two@example.com', 'farm-two')
        second_staff.office = 'New Active Farm'
        second_staff.save()
        unverified_staff = User.objects.create_user(
            email='farm-unverified@example.com',
            password='pass1234',
            username='farm-unverified',
            name='Unverified Farm User',
            office='New Active Farm',
            status=UserStatusChoices.ACTIVE,
        )
        inactive_staff = create_user(
            'farm-inactive@example.com',
            'farm-inactive',
            status=UserStatusChoices.INACTIVE,
        )
        inactive_staff.office = 'New Active Farm'
        inactive_staff.save()
        other_staff = create_user('other-farm@example.com', 'other-farm')
        other_staff.office = 'Other Farm'
        other_staff.save()
        self.client.force_authenticate(user=self.admin)

        response = self.client.post(
            '/api/data-requests/',
            {
                'recipient_scope': 'establishment',
                'establishments': ['New Active Farm'],
                'title': 'Farm data needed',
                'message': 'Please upload the latest visitor workbook.',
            },
            format='json',
        )

        recipient_ids = set(
            SubmissionNotification.objects.filter(
                notification_type=SubmissionNotification.TYPE_DATA_REQUEST
            ).values_list('recipient_id', flat=True)
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data['created_count'], 3)
        self.assertEqual(recipient_ids, {first_staff.pk, second_staff.pk, unverified_staff.pk})
        self.assertNotIn(inactive_staff.pk, recipient_ids)
        self.assertNotIn(other_staff.pk, recipient_ids)

    def test_user_cannot_send_data_request_notification(self):
        self.client.force_authenticate(user=self.staff)

        response = self.client.post(
            '/api/data-requests/',
            {
                'recipient_ids': [self.staff.pk],
                'title': 'Request',
                'message': 'Please upload a file.',
            },
            format='json',
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(SubmissionNotification.objects.filter(notification_type=SubmissionNotification.TYPE_DATA_REQUEST).exists())
