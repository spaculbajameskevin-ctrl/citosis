import csv
import json
import re
from datetime import datetime, timedelta
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from django.db.models.deletion import ProtectedError
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from rest_framework import status
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import xlrd

from accounts.models import User
from accounts.permissions import (
    CanApproveSubmissionPermission,
    CanEditDataPermission,
    IsActiveSystemUser,
    IsAdminDashboardUser,
    IsSubmissionWorkspaceAdmin,
    IsSuperAdminOnly,
)
from audit.models import ActivityLog, DataSubmission, RecycleBin, SubmissionNotification
from audit.serializers import ActivityLogSerializer, DataSubmissionSerializer, RecycleBinSerializer, SubmissionNotificationSerializer
from audit.services import (
    append_submission_review_event,
    build_submission_priority_metadata,
    create_admin_submission_alerts,
    create_data_request_notifications,
    create_submission_notification,
    import_backup_payload,
    log_action,
    purge_recycle_entry,
    restore_recycle_entry,
    serialize_backup_payload,
)
from citosis_pro.common import RecordCategoryChoices, UserStatusChoices, can_access_admin_dashboard
from tourism.models import TourismRecord, Visitor
from tourism.serializers import TourismRecordSerializer


HEADER_PATTERN = re.compile(r'[^a-z0-9]+')
EMAIL_PATTERN = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
VALID_RECORD_CATEGORIES = {value.lower() for value in RecordCategoryChoices.values}
EXCEL_VALIDATION_PROFILES = (
    {
        'key': 'visitor',
        'label': 'Visitor Sheet',
        'required_columns': {
            'Date': {'date', 'visit date', 'visit_date'},
            'Name': {'name', 'visitor name'},
            'Place': {'place', 'destination', 'destination name'},
            'Origin': {'origin', 'residence', 'source'},
        },
        'duplicate_columns': ['Date', 'Name', 'Place', 'Origin'],
    },
    {
        'key': 'record',
        'label': 'Tourism Record Sheet',
        'required_columns': {
            'Place': {'place', 'name', 'destination'},
            'Location': {'location', 'address'},
            'Category': {'category', 'type'},
        },
        'duplicate_columns': ['Place', 'Location'],
    },
)
EXCEL_VALIDATION_PROFILE_MAP = {profile['key']: profile for profile in EXCEL_VALIDATION_PROFILES}
EXCEL_TEMPLATE_DEFINITIONS = {
    'visitor': {
        'label': 'Visitor Sheet',
        'filename': 'visitor-upload-template.xlsx',
        'sheet_name': 'Visitors',
        'headers': ['Date', 'Name', 'Place', 'Origin'],
        'sample_rows': [
            ['2026-04-01', 'Ana Cruz', 'River Park', 'Malaybalay City'],
            ['2026-04-02', 'Ben Lopez', 'Bean Zone', 'Valencia City'],
        ],
        'notes': [
            'Use one row per visitor entry.',
            'Date accepts formats like YYYY-MM-DD or MM/DD/YYYY.',
            'Do not rename the header row if you want automatic validation to work best.',
        ],
    },
    'record': {
        'label': 'Tourism Record Sheet',
        'filename': 'DOT-Forms.xlsx',
        'file_path': Path('static') / 'templates' / 'DOT-Forms.xlsx',
        'sheet_name': 'Tourism Records',
        'headers': ['Place', 'Location', 'Category'],
        'sample_rows': [
            ['Bean Zone', 'Capitol Grounds', 'Event'],
            ['Capistrano', 'Binalbagan Simaya Malaybalay City', 'Beach'],
        ],
        'notes': [
            'Use one row per destination record.',
            'Category should match the allowed tourism categories used in the system.',
            'Keep the first header row intact for easier validation.',
        ],
    },
}
SUPPORTED_DATA_FILE_EXTENSIONS = {'.csv', '.xls', '.xlsx'}
SUPPORTED_DATA_FILE_EXTENSION_LABEL = '.csv, .xls, or .xlsx'


class DashboardOverviewView(APIView):
    permission_classes = [IsAdminDashboardUser]

    def get(self, request):
        today = timezone.localdate()
        submissions_today = DataSubmission.objects.exclude(status=DataSubmission.STATUS_DRAFT).filter(created_at__date=today)
        pending_submissions = DataSubmission.objects.filter(status=DataSubmission.STATUS_PENDING)
        in_review_submissions = DataSubmission.objects.filter(status=DataSubmission.STATUS_IN_REVIEW)
        completed_review_submissions = DataSubmission.objects.filter(
            status__in=[DataSubmission.STATUS_APPROVED, DataSubmission.STATUS_REJECTED]
        )

        records_count = TourismRecord.objects.count()
        visitors_count = Visitor.objects.count()
        users_count = User.objects.filter(status=UserStatusChoices.ACTIVE).count()
        recycle_count = RecycleBin.objects.filter(restored_at__isnull=True).count()

        popularity = list(
            Visitor.objects.values('place')
            .annotate(total=Count('id'))
            .order_by('-total', 'place')[:5]
        )

        recent_activity = ActivityLog.objects.select_related('user').order_by('-created_at')[:8]
        recent_records = TourismRecord.objects.select_related('created_by', 'updated_by').order_by('-updated_at')[:5]

        summary = [
            {
                'title': 'Checked In Visitors',
                'value': Visitor.objects.filter(status='Checked In').count(),
                'description': 'Guests currently marked as checked in.',
            },
            {
                'title': 'Checked Out Visitors',
                'value': Visitor.objects.filter(status='Checked Out').count(),
                'description': 'Guests who already completed their visit.',
            },
            {
                'title': 'Inactive Users',
                'value': User.objects.filter(status='Inactive').count(),
                'description': 'Staff accounts currently not active.',
            },
            {
                'title': 'Recent Logs',
                'value': ActivityLog.objects.count(),
                'description': 'Total activity log entries stored.',
            },
        ]

        rejection_count = completed_review_submissions.filter(status=DataSubmission.STATUS_REJECTED).count()
        completed_review_count = completed_review_submissions.count()
        rejection_rate = round((rejection_count / completed_review_count) * 100, 1) if completed_review_count else 0

        average_review_seconds = _calculate_average_review_seconds(completed_review_submissions)
        yesterday_submissions_count = (
            DataSubmission.objects.exclude(status=DataSubmission.STATUS_DRAFT)
            .filter(created_at__date=today - timedelta(days=1))
            .count()
        )
        pending_count = pending_submissions.count()
        in_review_count = in_review_submissions.count()

        dashboard_metrics = [
            {
                'key': 'submissions_today',
                'title': 'Submissions Today',
                'value': submissions_today.count(),
                'display_value': str(submissions_today.count()),
                'description': 'New CSV and Excel files sent to the admin queue since midnight.',
                'meta': (
                    f'{yesterday_submissions_count} yesterday'
                    if yesterday_submissions_count
                    else 'No submissions were received yesterday.'
                ),
                'tone': 'primary',
            },
            {
                'key': 'pending_reviews',
                'title': 'Pending Reviews',
                'value': pending_count + in_review_count,
                'display_value': str(pending_count + in_review_count),
                'description': 'Files still waiting for an admin decision.',
                'meta': (
                    f'{in_review_count} already in review, {pending_count} still pending.'
                    if pending_count + in_review_count
                    else 'No submissions are waiting in the review queue.'
                ),
                'tone': 'warning',
            },
            {
                'key': 'rejection_rate',
                'title': 'Rejection Rate',
                'value': rejection_rate,
                'display_value': _format_dashboard_percentage(rejection_rate),
                'description': 'Rejected out of completed admin review decisions.',
                'meta': (
                    f'{rejection_count} rejected out of {completed_review_count} reviewed.'
                    if completed_review_count
                    else 'No completed review decisions yet.'
                ),
                'tone': 'danger' if rejection_rate >= 40 else 'neutral',
            },
            {
                'key': 'avg_review_time',
                'title': 'Avg Review Time',
                'value': average_review_seconds,
                'display_value': _format_dashboard_duration(average_review_seconds),
                'description': 'Average time from submission to final review decision.',
                'meta': (
                    f'Based on {completed_review_count} approved or rejected file(s).'
                    if completed_review_count
                    else 'Average review time will appear after the first completed decision.'
                ),
                'tone': 'success',
            },
        ]

        dashboard_charts = {
            'submissions': _build_dashboard_activity_chart(
                queryset=DataSubmission.objects.exclude(status=DataSubmission.STATUS_DRAFT),
                field_name='created_at',
                title='Submissions over time',
                summary_template='{current_total} submission(s) received in the last 7 days.',
                zero_summary='No submissions were received in the last 7 days.',
                period_days=7,
            ),
            'visitors': _build_dashboard_growth_chart(
                queryset=Visitor.objects.all(),
                field_name='created_at',
                title='Visitors growth',
                period_days=7,
            ),
            'destinations': _build_destination_trend_chart(),
        }

        return Response(
            {
                'stats': {
                    'records': records_count,
                    'visitors': visitors_count,
                    'users': users_count,
                    'recycle': recycle_count,
                },
                'intelligence': {
                    'metrics': dashboard_metrics,
                    'charts': dashboard_charts,
                },
                'destination_popularity': popularity,
                'recent_activity': ActivityLogSerializer(recent_activity, many=True).data,
                'recent_records': TourismRecordSerializer(recent_records, many=True, context={'request': request}).data,
                'summary': summary,
            }
        )


def _build_dashboard_activity_chart(queryset, field_name, title, summary_template, zero_summary, period_days=7):
    current_points, current_total = _build_daily_series(queryset, field_name=field_name, period_days=period_days)
    _, previous_total = _build_daily_series(
        queryset,
        field_name=field_name,
        period_days=period_days,
        end_date=timezone.localdate() - timedelta(days=period_days),
    )
    return {
        'title': title,
        'window_label': f'Last {period_days} days',
        'summary': summary_template.format(current_total=current_total) if current_total else zero_summary,
        'change_label': _build_window_comparison_label(current_total, previous_total, 'submission'),
        'points': current_points,
        'total': current_total,
    }


def _build_dashboard_growth_chart(queryset, field_name, title, period_days=7):
    today = timezone.localdate()
    start_date = today - timedelta(days=period_days - 1)
    daily_points, added_this_window = _build_daily_series(
        queryset,
        field_name=field_name,
        period_days=period_days,
        end_date=today,
    )
    baseline_total = 0
    for timestamp in queryset.values_list(field_name, flat=True):
        if not timestamp:
            continue
        local_date = _to_local_date(timestamp)
        if local_date and local_date < start_date:
            baseline_total += 1

    running_total = baseline_total
    cumulative_points = []
    for point in daily_points:
        running_total += point['value']
        cumulative_points.append({**point, 'value': running_total})

    return {
        'title': title,
        'window_label': f'Last {period_days} days',
        'summary': (
            f'{added_this_window} new visitor entr{"y" if added_this_window == 1 else "ies"} added in the last {period_days} days.'
            if added_this_window
            else f'No new visitor entries were added in the last {period_days} days.'
        ),
        'change_label': (
            f'{running_total} total visitor entr{"y" if running_total == 1 else "ies"} tracked.'
            if running_total
            else 'No visitor entries tracked yet.'
        ),
        'points': cumulative_points,
        'total': running_total,
        'delta': added_this_window,
    }


def _build_destination_trend_chart(period_days=7):
    today = timezone.localdate()
    recent_start = today - timedelta(days=29)

    recent_places = list(
        Visitor.objects.exclude(place='')
        .filter(created_at__date__gte=recent_start)
        .values('place')
        .annotate(total=Count('id'))
        .order_by('-total', 'place')[:5]
    )
    if not recent_places:
        recent_places = list(
            Visitor.objects.exclude(place='')
            .values('place')
            .annotate(total=Count('id'))
            .order_by('-total', 'place')[:5]
        )

    items = []
    for row in recent_places:
        place_name = row.get('place') or 'Unknown destination'
        place_queryset = Visitor.objects.filter(place=place_name)
        trend_points, weekly_total = _build_daily_series(
            place_queryset,
            field_name='created_at',
            period_days=period_days,
            end_date=today,
        )
        items.append(
            {
                'label': place_name,
                'total': row.get('total', 0),
                'recent_total': weekly_total,
                'summary': (
                    f'{weekly_total} visit entr{"y" if weekly_total == 1 else "ies"} in the last {period_days} days.'
                    if weekly_total
                    else f'No recent visitor activity in the last {period_days} days.'
                ),
                'points': trend_points,
            }
        )

    return {
        'title': 'Popular destinations trends',
        'window_label': f'Last {period_days} days',
        'summary': 'Recent visitor momentum for the most popular destinations.',
        'items': items,
    }


def _build_daily_series(queryset, field_name, period_days=7, end_date=None):
    resolved_end_date = end_date or timezone.localdate()
    start_date = resolved_end_date - timedelta(days=period_days - 1)
    counts_by_date = {
        start_date + timedelta(days=offset): 0
        for offset in range(period_days)
    }

    for timestamp in queryset.values_list(field_name, flat=True):
        if not timestamp:
            continue
        local_date = _to_local_date(timestamp)
        if local_date in counts_by_date:
            counts_by_date[local_date] += 1

    points = []
    total = 0
    for offset in range(period_days):
        current_date = start_date + timedelta(days=offset)
        value = counts_by_date[current_date]
        total += value
        points.append(
            {
                'date': current_date.isoformat(),
                'label': current_date.strftime('%b %d'),
                'value': value,
            }
        )
    return points, total


def _to_local_date(value):
    if value is None:
        return None
    if hasattr(value, 'date'):
        try:
            if timezone.is_aware(value):
                return timezone.localtime(value).date()
            return value.date()
        except (TypeError, ValueError):
            return None
    return None


def _calculate_average_review_seconds(queryset):
    durations = []
    for created_at, updated_at in queryset.values_list('created_at', 'updated_at'):
        if not created_at or not updated_at:
            continue
        duration_seconds = max(0, int((updated_at - created_at).total_seconds()))
        durations.append(duration_seconds)
    if not durations:
        return 0
    return round(sum(durations) / len(durations))


def _format_dashboard_percentage(value):
    percentage = round(float(value or 0), 1)
    if percentage.is_integer():
        return f'{int(percentage)}%'
    return f'{percentage:.1f}%'


def _format_dashboard_duration(seconds):
    total_seconds = max(0, int(seconds or 0))
    total_minutes = round(total_seconds / 60)
    if total_minutes <= 0:
        return '0m'
    if total_minutes < 60:
        return f'{total_minutes}m'
    hours, minutes = divmod(total_minutes, 60)
    return f'{hours}h' if not minutes else f'{hours}h {minutes}m'


def _build_window_comparison_label(current_total, previous_total, singular_label):
    current_total = int(current_total or 0)
    previous_total = int(previous_total or 0)
    plural_label = singular_label if current_total == 1 else f'{singular_label}s'

    if previous_total == current_total:
        return f'Steady at {current_total} {plural_label} versus the previous window.'
    if previous_total == 0:
        return (
            f'Up from 0 in the previous window.'
            if current_total
            else f'No {plural_label} in either window.'
        )

    difference = current_total - previous_total
    direction = 'up' if difference > 0 else 'down'
    magnitude = abs(difference)
    change_percentage = round((magnitude / previous_total) * 100)
    comparison_label = singular_label if magnitude == 1 else f'{singular_label}s'
    return f'{direction.title()} {change_percentage}% ({magnitude} {comparison_label}) versus the previous window.'


class DashboardExcelPreviewView(APIView):
    permission_classes = [CanEditDataPermission]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'detail': 'Choose a CSV or Excel file to preview first.'}, status=status.HTTP_400_BAD_REQUEST)

        extension = Path(file.name or '').suffix.lower()
        if extension not in SUPPORTED_DATA_FILE_EXTENSIONS:
            return Response({'detail': f'Invalid file format. Please upload a {SUPPORTED_DATA_FILE_EXTENSION_LABEL} file only.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            preview = build_excel_preview_payload(file, extension)
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception:
            return Response(
                {'detail': 'This CSV or Excel file could not be read. It may be corrupted, password-protected, or use an unsupported structure.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        log_action(
            request.user,
            'Previewed data file',
            f'Previewed data file "{file.name}" in the admin dashboard.',
            request,
        )
        return Response(preview)


class DataSubmissionPreviewView(APIView):
    permission_classes = [IsActiveSystemUser]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'detail': 'Choose a CSV or Excel file to preview first.'}, status=status.HTTP_400_BAD_REQUEST)

        extension = Path(file.name or '').suffix.lower()
        if extension not in SUPPORTED_DATA_FILE_EXTENSIONS:
            return Response({'detail': f'Invalid file format. Please upload a {SUPPORTED_DATA_FILE_EXTENSION_LABEL} file only.'}, status=status.HTTP_400_BAD_REQUEST)

        selected_sheet_name = str(request.data.get('sheet_name', '')).strip() or None
        try:
            mapping_profile_key = _parse_mapping_profile_key(request.data.get('mapping_profile_key'))
            column_mapping = _parse_submission_column_mapping(request.data.get('column_mapping'))
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        try:
            preview = build_excel_preview_payload(
                file,
                extension,
                sheet_name=selected_sheet_name,
                max_preview_rows=100,
                max_preview_columns=40,
                mapping_profile_key=mapping_profile_key,
                column_mapping=column_mapping,
            )
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception:
            return Response(
                {'detail': 'This CSV or Excel file could not be read. It may be corrupted, password-protected, or use an unsupported structure.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        _attach_submission_version_context(preview, request.user, file.name)
        return Response(preview)


class DataSubmissionTemplateDownloadView(APIView):
    permission_classes = [IsActiveSystemUser]

    def get(self, request, template_key):
        template_definition = EXCEL_TEMPLATE_DEFINITIONS.get(template_key)
        if not template_definition:
            return Response({'detail': 'Template not found.'}, status=status.HTTP_404_NOT_FOUND)

        file_path = template_definition.get('file_path')
        if file_path:
            resolved_file_path = Path(__file__).resolve().parent.parent / file_path
            if not resolved_file_path.exists():
                return Response({'detail': 'Template file not found.'}, status=status.HTTP_404_NOT_FOUND)
            response = HttpResponse(
                resolved_file_path.read_bytes(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            response['Content-Disposition'] = f'attachment; filename="{template_definition["filename"]}"'
            return response

        workbook = _build_submission_template_workbook(template_definition)
        response = HttpResponse(
            workbook.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{template_definition["filename"]}"'
        return response


class DataSubmissionListCreateView(APIView):
    permission_classes = [IsActiveSystemUser]
    parser_classes = [MultiPartParser, FormParser]

    def get(self, request):
        queryset = DataSubmission.objects.select_related('submitted_by', 'last_reviewed_by')
        if not can_access_admin_dashboard(request.user):
            queryset = queryset.filter(submitted_by=request.user)
        else:
            queryset = queryset.exclude(status=DataSubmission.STATUS_DRAFT)
        serializer = DataSubmissionSerializer(queryset.order_by('-updated_at', '-created_at', '-id'), many=True, context={'request': request})
        return Response(serializer.data)

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'detail': 'Choose a CSV or Excel file to send first.'}, status=status.HTTP_400_BAD_REQUEST)

        extension = Path(file.name or '').suffix.lower()
        if extension not in SUPPORTED_DATA_FILE_EXTENSIONS:
            return Response({'detail': f'Invalid file format. Please upload a {SUPPORTED_DATA_FILE_EXTENSION_LABEL} file only.'}, status=status.HTTP_400_BAD_REQUEST)

        selected_sheet_name = str(request.data.get('sheet_name', '')).strip() or None
        save_mode = str(request.data.get('save_mode', 'send')).strip().lower() or 'send'
        if save_mode not in {'send', 'draft'}:
            return Response({'detail': 'Invalid save mode. Choose either "send" or "draft".'}, status=status.HTTP_400_BAD_REQUEST)
        submission_notes = str(request.data.get('notes', '')).strip()
        try:
            mapping_profile_key = _parse_mapping_profile_key(request.data.get('mapping_profile_key'))
            column_mapping = _parse_submission_column_mapping(request.data.get('column_mapping'))
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        try:
            preview = build_excel_preview_payload(
                file,
                extension,
                sheet_name=selected_sheet_name,
                max_preview_rows=100,
                max_preview_columns=40,
                mapping_profile_key=mapping_profile_key,
                column_mapping=column_mapping,
            )
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception:
            return Response(
                {'detail': 'This CSV or Excel file could not be read. It may be corrupted, password-protected, or use an unsupported structure.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        file.seek(0)
        version_number = _get_next_data_submission_version(request.user, file.name, preview['sheet_name'])
        next_status = DataSubmission.STATUS_DRAFT if save_mode == 'draft' else DataSubmission.STATUS_PENDING
        submission = DataSubmission.objects.create(
            submitted_by=request.user,
            file=file,
            original_filename=file.name,
            sheet_name=preview['sheet_name'],
            headers=preview['headers'],
            preview_rows=preview['rows'],
            total_rows=preview['total_rows'],
            total_columns=preview['total_columns'],
            mapping_profile_key=preview.get('mapping_profile_key', ''),
            column_mapping=preview.get('column_mapping', {}),
            submission_notes=submission_notes,
            version_number=version_number,
            status=next_status,
        )
        append_submission_review_event(
            submission,
            actor=request.user,
            action_key='draft_saved' if next_status == DataSubmission.STATUS_DRAFT else 'submitted',
            title='Draft saved' if next_status == DataSubmission.STATUS_DRAFT else 'Submission sent',
            summary=(
                f'Saved version v{version_number} as a draft for later review.'
                if next_status == DataSubmission.STATUS_DRAFT
                else f'Sent version v{version_number} to the admin review queue.'
            ),
            details={
                'status': next_status,
                'version_label': f'v{version_number}',
                'sheet_name': preview['sheet_name'],
                'file_name': file.name,
            },
            mark_reviewed=False,
        )
        submission.save(update_fields=['review_history', 'updated_at'])
        if next_status != DataSubmission.STATUS_DRAFT:
            create_admin_submission_alerts(submission, actor=request.user)
        log_action(
            request.user,
            'Saved data draft' if next_status == DataSubmission.STATUS_DRAFT else 'Sent data file',
            (
                f'Saved data file "{file.name}" (sheet "{preview["sheet_name"]}") as draft version v{version_number}.'
                if next_status == DataSubmission.STATUS_DRAFT
                else f'Sent data file "{file.name}" (sheet "{preview["sheet_name"]}") to the admin dashboard as version v{version_number}.'
            ),
            request,
        )
        serializer = DataSubmissionSerializer(submission, context={'request': request})
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class DataSubmissionSendView(APIView):
    permission_classes = [IsActiveSystemUser]
    parser_classes = [JSONParser, FormParser, MultiPartParser]

    def post(self, request, pk):
        submission = DataSubmission.objects.select_related('submitted_by', 'last_reviewed_by').filter(pk=pk, submitted_by=request.user).first()
        if not submission:
            return Response({'detail': 'Draft submission not found.'}, status=status.HTTP_404_NOT_FOUND)

        if submission.status != DataSubmission.STATUS_DRAFT:
            return Response({'detail': 'Only draft submissions can be sent later.'}, status=status.HTTP_400_BAD_REQUEST)

        submission_notes = request.data.get('notes')
        if submission_notes is not None:
            submission.submission_notes = str(submission_notes).strip()

        submission.status = DataSubmission.STATUS_PENDING
        append_submission_review_event(
            submission,
            actor=request.user,
            action_key='submitted',
            title='Submission sent',
            summary=f'Sent draft version v{submission.version_number} to the admin review queue.',
            details={
                'status': DataSubmission.STATUS_PENDING,
                'version_label': f'v{submission.version_number}',
                'sheet_name': submission.sheet_name,
                'file_name': submission.original_filename,
            },
            mark_reviewed=False,
        )
        submission.save(update_fields=['status', 'submission_notes', 'review_history', 'updated_at'])
        create_admin_submission_alerts(submission, actor=request.user)

        log_action(
            request.user,
            'Sent saved data draft',
            f'Sent draft data file "{submission.original_filename}" (sheet "{submission.sheet_name}") version v{submission.version_number} to the admin dashboard.',
            request,
        )

        serializer = DataSubmissionSerializer(submission, context={'request': request})
        return Response(serializer.data)


class DataSubmissionStatusUpdateView(APIView):
    permission_classes = [CanApproveSubmissionPermission]

    def patch(self, request, pk):
        submission = DataSubmission.objects.select_related('submitted_by', 'last_reviewed_by').filter(pk=pk).first()
        if not submission:
            return Response({'detail': 'Data submission not found.'}, status=status.HTTP_404_NOT_FOUND)

        next_status = str(request.data.get('status', '')).strip() or submission.status
        admin_feedback = str(request.data.get('admin_feedback', submission.admin_feedback or '')).strip()
        try:
            _apply_submission_status_update(submission, next_status, admin_feedback, actor=request.user, request=request)
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        serializer = DataSubmissionSerializer(submission, context={'request': request})
        return Response(serializer.data)


class DataSubmissionBulkStatusUpdateView(APIView):
    permission_classes = [CanApproveSubmissionPermission]
    parser_classes = [JSONParser, FormParser]

    def post(self, request):
        try:
            submission_ids = _parse_submission_id_list(request.data.get('submission_ids'))
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        next_status = str(request.data.get('status', '')).strip()
        admin_feedback = str(request.data.get('admin_feedback', '')).strip()

        submissions = list(
            DataSubmission.objects.select_related('submitted_by', 'last_reviewed_by')
            .filter(pk__in=submission_ids)
            .exclude(status=DataSubmission.STATUS_DRAFT)
        )
        found_ids = {submission.pk for submission in submissions}
        missing_ids = [submission_id for submission_id in submission_ids if submission_id not in found_ids]
        if missing_ids:
            return Response({'detail': 'One or more selected submissions could not be found.'}, status=status.HTTP_404_NOT_FOUND)

        updated_submissions = []
        for submission in submissions:
            try:
                changed = _apply_submission_status_update(
                    submission,
                    next_status,
                    admin_feedback,
                    actor=request.user,
                    request=request,
                    log_action_name='Updated Excel submissions in bulk',
                )
            except ValueError as error:
                return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)
            if changed:
                updated_submissions.append(submission)
            else:
                updated_submissions.append(submission)

        serializer = DataSubmissionSerializer(updated_submissions, many=True, context={'request': request})
        return Response(
            {
                'updated_count': len(updated_submissions),
                'status': next_status,
                'updated_submissions': serializer.data,
            }
        )


class DataSubmissionBulkExportView(APIView):
    permission_classes = [IsAdminDashboardUser]

    def get(self, request):
        try:
            submission_ids = _parse_submission_id_list(request.query_params.get('ids'))
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        submissions = list(
            DataSubmission.objects.select_related('submitted_by')
            .filter(pk__in=submission_ids)
            .exclude(status=DataSubmission.STATUS_DRAFT)
            .order_by('-updated_at', '-created_at', '-id')
        )
        found_ids = {submission.pk for submission in submissions}
        missing_ids = [submission_id for submission_id in submission_ids if submission_id not in found_ids]
        if missing_ids:
            return Response({'detail': 'One or more selected submissions could not be found.'}, status=status.HTTP_404_NOT_FOUND)

        zip_buffer = BytesIO()
        used_names = set()
        exported_count = 0
        manifest_rows = []
        with ZipFile(zip_buffer, 'w', compression=ZIP_DEFLATED) as zip_file:
            for submission in submissions:
                if not submission.file:
                    continue

                base_name = Path(submission.original_filename or submission.file.name).name or f'submission-{submission.pk}.xlsx'
                safe_name = _build_unique_export_filename(
                    used_names,
                    f'{submission.version_number:02d}-submission-{submission.pk}-{base_name}',
                )
                with submission.file.open('rb') as source_file:
                    zip_file.writestr(safe_name, source_file.read())

                manifest_rows.append(
                    {
                        'id': submission.pk,
                        'filename': submission.original_filename,
                        'sheet_name': submission.sheet_name,
                        'status': submission.status,
                        'submitted_by': submission.submitted_by.email,
                        'version': f'v{submission.version_number}',
                    }
                )
                exported_count += 1

            zip_file.writestr('manifest.json', json.dumps(manifest_rows, indent=2))

        if not exported_count:
            return Response({'detail': 'None of the selected submissions have downloadable files.'}, status=status.HTTP_400_BAD_REQUEST)

        zip_buffer.seek(0)
        timestamp = timezone.now().strftime('%Y%m%d-%H%M%S')
        response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="incoming-data-export-{timestamp}.zip"'

        log_action(
            request.user,
            'Exported Excel submissions in bulk',
            f'Exported {exported_count} Excel submission file(s) from the incoming data queue.',
            request,
        )
        return response


class DataSubmissionWorkspaceView(APIView):
    permission_classes = [IsSubmissionWorkspaceAdmin]
    parser_classes = [JSONParser, FormParser]

    def get(self, request, pk):
        submission = DataSubmission.objects.select_related('submitted_by', 'last_reviewed_by').filter(pk=pk).first()
        if not submission:
            return Response({'detail': 'Data submission not found.'}, status=status.HTTP_404_NOT_FOUND)

        return Response(_build_submission_workspace_response(submission, request))

    def patch(self, request, pk):
        submission = DataSubmission.objects.select_related('submitted_by', 'last_reviewed_by').filter(pk=pk).first()
        if not submission:
            return Response({'detail': 'Data submission not found.'}, status=status.HTTP_404_NOT_FOUND)

        try:
            workspace_payload = _parse_submission_workspace_request_payload(request.data)
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        workspace_preview = _resolve_submission_workspace_payload(
            submission,
            headers=workspace_payload['headers'],
            preview_rows=workspace_payload['preview_rows'],
            mapping_profile_key=workspace_payload['mapping_profile_key'],
            column_mapping=workspace_payload['column_mapping'],
        )
        workspace_diff = _build_submission_content_diff(
            submission.headers,
            submission.preview_rows,
            workspace_preview['headers'],
            workspace_preview['rows'],
            max_items=10,
        )
        previous_mapping_profile_key = submission.mapping_profile_key or ''
        mapping_profile_changed = (submission.mapping_profile_key or '') != (workspace_preview['mapping_profile_key'] or '')
        previous_column_mapping = submission.column_mapping or {}
        next_column_mapping = workspace_preview['column_mapping'] or {}
        mapping_changes = sorted(
            field_name
            for field_name in set(previous_column_mapping) | set(next_column_mapping)
            if previous_column_mapping.get(field_name) != next_column_mapping.get(field_name)
        )

        submission.headers = workspace_preview['headers']
        submission.preview_rows = workspace_preview['rows']
        submission.total_rows = workspace_preview['total_rows']
        submission.total_columns = workspace_preview['total_columns']
        submission.mapping_profile_key = workspace_preview['mapping_profile_key']
        submission.column_mapping = workspace_preview['column_mapping']
        append_submission_review_event(
            submission,
            actor=request.user,
            action_key='workspace_saved',
            title='Preview workspace updated',
            summary=_build_submission_workspace_update_summary(
                workspace_diff,
                mapping_profile_changed=mapping_profile_changed,
                mapping_change_count=len(mapping_changes),
            ),
            details={
                'comparison_summary': workspace_diff['summary'],
                'comparison_metrics': workspace_diff['metrics'],
                'changes': workspace_diff['items'],
                'mapping_profile_before': previous_mapping_profile_key,
                'mapping_profile_after': workspace_preview['mapping_profile_key'] or '',
                'mapping_changes': [
                    {
                        'field_name': field_name,
                        'before': previous_column_mapping.get(field_name, ''),
                        'after': next_column_mapping.get(field_name, ''),
                    }
                    for field_name in mapping_changes
                ],
            },
        )
        submission.save(
            update_fields=[
                'headers',
                'preview_rows',
                'total_rows',
                'total_columns',
                'mapping_profile_key',
                'column_mapping',
                'review_history',
                'last_reviewed_by',
                'last_reviewed_at',
                'updated_at',
            ]
        )

        log_action(
            request.user,
            'Updated submission preview workspace',
            f'Updated preview data for data file "{submission.original_filename}" (sheet "{submission.sheet_name}").',
            request,
        )

        return Response(_build_submission_workspace_response(submission, request, workspace_preview=workspace_preview))


class DataSubmissionWorkspaceValidateView(APIView):
    permission_classes = [CanEditDataPermission]
    parser_classes = [JSONParser, FormParser]

    def post(self, request, pk):
        submission = DataSubmission.objects.select_related('submitted_by', 'last_reviewed_by').filter(pk=pk).first()
        if not submission:
            return Response({'detail': 'Data submission not found.'}, status=status.HTTP_404_NOT_FOUND)

        try:
            workspace_payload = _parse_submission_workspace_request_payload(request.data)
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        workspace_preview = _resolve_submission_workspace_payload(
            submission,
            headers=workspace_payload['headers'],
            preview_rows=workspace_payload['preview_rows'],
            mapping_profile_key=workspace_payload['mapping_profile_key'],
            column_mapping=workspace_payload['column_mapping'],
        )

        return Response(_build_submission_workspace_response(submission, request, workspace_preview=workspace_preview))


class DataRequestTargetListView(APIView):
    permission_classes = [CanApproveSubmissionPermission]

    def get(self, request):
        targets = [
            user
            for user in _get_data_request_target_queryset()
            if user.pk != request.user.pk and not can_access_admin_dashboard(user)
        ]
        return Response([_serialize_data_request_target(user) for user in targets])


class DataRequestCreateView(APIView):
    permission_classes = [CanApproveSubmissionPermission]
    parser_classes = [JSONParser, FormParser]

    def post(self, request):
        title = str(request.data.get('title', '')).strip() or 'Data request from admin'
        message = str(request.data.get('message', '')).strip()
        if not message:
            return Response({'detail': 'Describe the data you need from the user.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            due_date = _parse_data_request_due_date(request.data.get('due_date'))
            recipients = _resolve_data_request_recipients(request)
        except ValueError as error:
            return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)

        if not recipients:
            return Response({'detail': 'Choose at least one active user to notify.'}, status=status.HTTP_400_BAD_REQUEST)

        notifications = create_data_request_notifications(
            recipients,
            actor=request.user,
            title=title,
            message=message,
            due_date=due_date,
        )
        recipient_labels = ', '.join(user.email for user in recipients[:6])
        if len(recipients) > 6:
            recipient_labels += f', and {len(recipients) - 6} more'
        log_action(
            request.user,
            'Requested data from users',
            f'Sent "{title}" data request to {len(recipients)} user(s): {recipient_labels}.',
            request,
        )
        serializer = SubmissionNotificationSerializer(notifications, many=True)
        return Response(
            {
                'detail': f'Data request sent to {len(notifications)} user(s).',
                'created_count': len(notifications),
                'notifications': serializer.data,
            },
            status=status.HTTP_201_CREATED,
        )


class DataRequestHistoryView(APIView):
    permission_classes = [CanApproveSubmissionPermission]

    def get(self, request):
        try:
            limit = int(request.query_params.get('limit', '50'))
        except (TypeError, ValueError):
            limit = 50
        limit = max(1, min(limit, 100))
        queryset = (
            SubmissionNotification.objects.select_related('recipient', 'created_by')
            .filter(notification_type=SubmissionNotification.TYPE_DATA_REQUEST)
            .order_by('-created_at', '-id')
        )
        serializer = SubmissionNotificationSerializer(queryset[:limit], many=True)
        return Response(serializer.data)


class SubmissionNotificationListView(APIView):
    permission_classes = [IsActiveSystemUser]

    def get(self, request):
        try:
            limit = int(request.query_params.get('limit', '12'))
        except (TypeError, ValueError):
            limit = 12
        limit = max(1, min(limit, 50))
        queryset = SubmissionNotification.objects.select_related('submission', 'created_by').filter(recipient=request.user).order_by('is_read', '-created_at', '-id')
        serializer = SubmissionNotificationSerializer(queryset[:limit], many=True)
        return Response(serializer.data)


class SubmissionNotificationReadView(APIView):
    permission_classes = [IsActiveSystemUser]

    def post(self, request, pk):
        notification = SubmissionNotification.objects.select_related('submission', 'created_by').filter(pk=pk, recipient=request.user).first()
        if not notification:
            return Response({'detail': 'Notification not found.'}, status=status.HTTP_404_NOT_FOUND)

        if not notification.is_read:
            notification.is_read = True
            notification.read_at = timezone.now()
            notification.save(update_fields=['is_read', 'read_at'])

        serializer = SubmissionNotificationSerializer(notification)
        return Response(serializer.data)


class SubmissionNotificationReadAllView(APIView):
    permission_classes = [IsActiveSystemUser]

    def post(self, request):
        now = timezone.now()
        unread_queryset = SubmissionNotification.objects.filter(recipient=request.user, is_read=False)
        unread_count = unread_queryset.count()
        unread_queryset.update(is_read=True, read_at=now)
        return Response({'detail': 'Notifications marked as read.', 'updated_count': unread_count})


class ActivityLogListView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def get(self, request):
        search = request.query_params.get('search', '').strip()
        limit = int(request.query_params.get('limit', '25'))
        queryset = ActivityLog.objects.select_related('user').order_by('-created_at')
        if search:
            queryset = queryset.filter(action__icontains=search) | queryset.filter(details__icontains=search)
            queryset = queryset.distinct()
        queryset = queryset[:limit]
        return Response(ActivityLogSerializer(queryset, many=True).data)


class ActivityLogClearView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def delete(self, request):
        deleted_count, _ = ActivityLog.objects.all().delete()
        log_action(request.user, 'Cleared activity logs', f'Cleared {deleted_count} log rows.', request)
        return Response({'detail': 'Activity logs cleared.'})


class ActivityLogExportView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def get(self, request):
        payload = ActivityLogSerializer(ActivityLog.objects.select_related('user').order_by('-created_at'), many=True).data
        response = HttpResponse(json.dumps(payload, indent=2), content_type='application/json')
        response['Content-Disposition'] = 'attachment; filename="activity-logs.json"'
        log_action(request.user, 'Exported activity logs', 'Exported activity logs as JSON.', request)
        return response


class RecycleBinListView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def get(self, request):
        search = request.query_params.get('search', '').strip().lower()
        item_type = request.query_params.get('type', '').strip().lower()
        queryset = RecycleBin.objects.filter(restored_at__isnull=True).select_related('deleted_by', 'restored_by').order_by('-deleted_at')
        if item_type:
            queryset = queryset.filter(item_type=item_type)
        if search:
            filtered_ids = []
            for item in queryset:
                haystack = json.dumps(item.item_data).lower()
                if search in haystack or search in item.item_type.lower() or search in item.item_id.lower():
                    filtered_ids.append(item.pk)
            queryset = queryset.filter(pk__in=filtered_ids)
        return Response(RecycleBinSerializer(queryset, many=True).data)


class RecycleBinRestoreView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def post(self, request, pk):
        entry = RecycleBin.objects.filter(pk=pk, restored_at__isnull=True).first()
        if not entry:
            return Response({'detail': 'Recycle bin item not found.'}, status=status.HTTP_404_NOT_FOUND)
        restore_recycle_entry(entry, request.user, request)
        return Response({'detail': 'Item restored successfully.'})


class RecycleBinPurgeView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def delete(self, request, pk):
        entry = RecycleBin.objects.filter(pk=pk, restored_at__isnull=True).first()
        if not entry:
            return Response({'detail': 'Recycle bin item not found.'}, status=status.HTTP_404_NOT_FOUND)
        try:
            purge_recycle_entry(entry, request.user, request)
        except ProtectedError:
            return Response(
                {'detail': 'This item cannot be purged yet because other records still depend on it.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        return Response({'detail': 'Item permanently removed.'})


class RecycleBinEmptyView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def delete(self, request):
        entries = list(RecycleBin.objects.filter(restored_at__isnull=True))
        try:
            for entry in entries:
                purge_recycle_entry(entry, request.user, request)
        except ProtectedError:
            return Response(
                {'detail': 'One or more recycle bin items still have dependent records and could not be purged.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        return Response({'detail': 'Recycle bin emptied successfully.'})


class SystemBackupView(APIView):
    permission_classes = [IsSuperAdminOnly]

    def get(self, request):
        payload = serialize_backup_payload()
        response = HttpResponse(json.dumps(payload, indent=2), content_type='application/json')
        response['Content-Disposition'] = 'attachment; filename="citosis-pro-backup.json"'
        log_action(request.user, 'Exported backup', 'Generated a full JSON backup export.', request)
        return response


class SystemImportView(APIView):
    permission_classes = [IsSuperAdminOnly]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request):
        file = request.FILES.get('file')
        if file:
            payload = json.loads(file.read().decode('utf-8'))
        else:
            payload = request.data if isinstance(request.data, dict) else {}
        import_backup_payload(payload, actor=request.user, request=request)
        return Response({'detail': 'Backup imported successfully.'})


def build_excel_preview_payload(
    uploaded_file,
    extension,
    sheet_name=None,
    max_preview_rows=1000,
    max_preview_columns=60,
    mapping_profile_key=None,
    column_mapping=None,
):
    workbook_data = _extract_excel_workbook_data(uploaded_file, extension)
    selected_sheet = _select_excel_sheet(workbook_data, sheet_name)
    parsed_sheet = _parse_sheet_rows(selected_sheet, max_preview_columns=max_preview_columns)

    preview = _serialize_sheet_preview(
        file_name=uploaded_file.name,
        sheet_name=selected_sheet['name'],
        parsed_sheet=parsed_sheet,
        max_preview_rows=max_preview_rows,
    )
    preview['selected_sheet_name'] = selected_sheet['name']
    preview['sheets'] = [
        _summarize_excel_sheet(sheet_data, selected_sheet['name'], max_preview_columns)
        for sheet_data in workbook_data
    ]
    preview['validation'] = _build_sheet_validation_feedback(
        parsed_sheet['headers'],
        parsed_sheet['cleaned_rows'],
        mapping_profile_key=mapping_profile_key,
        column_mapping=column_mapping,
    )
    preview['mapping_profile_key'] = preview['validation'].get('profile_key') if preview['validation'].get('profile_key') != 'generic' else ''
    preview['column_mapping'] = preview['validation'].get('applied_column_mapping', {})
    return preview


def _extract_excel_workbook_data(uploaded_file, extension):
    if extension == '.csv':
        sheets = _extract_csv_data(uploaded_file)
    elif extension == '.xlsx':
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        try:
            sheets = []
            for worksheet in workbook.worksheets:
                raw_rows = []
                total_columns = 0
                for row in worksheet.iter_rows(values_only=True):
                    normalized = [_normalize_excel_value(value) for value in row]
                    total_columns = max(total_columns, len(normalized))
                    raw_rows.append(normalized)
                sheets.append(
                    {
                        'name': worksheet.title,
                        'raw_rows': raw_rows,
                        'total_columns': total_columns,
                    }
                )
        finally:
            workbook.close()
    else:
        file_bytes = uploaded_file.read()
        try:
            workbook = xlrd.open_workbook(file_contents=file_bytes)
        except xlrd.XLRDError as error:
            raise ValueError('This .xls file could not be opened. It may be corrupted or use an unsupported Excel format.') from error
        sheets = []
        for worksheet in workbook.sheets():
            raw_rows = []
            for row_index in range(worksheet.nrows):
                raw_rows.append([_normalize_excel_value(worksheet.cell_value(row_index, col_index)) for col_index in range(worksheet.ncols)])
            sheets.append(
                {
                    'name': worksheet.name,
                    'raw_rows': raw_rows,
                    'total_columns': worksheet.ncols,
                }
            )

    if hasattr(uploaded_file, 'seek'):
        uploaded_file.seek(0)

    if not sheets:
        raise ValueError('The uploaded file does not contain any readable table data.')

    return sheets


def _extract_csv_data(uploaded_file):
    file_bytes = uploaded_file.read()
    if isinstance(file_bytes, str):
        text_content = file_bytes
    else:
        text_content = _decode_csv_content(file_bytes)

    try:
        rows = [
            [_normalize_excel_value(cell) for cell in row]
            for row in csv.reader(StringIO(text_content))
        ]
    except csv.Error as error:
        raise ValueError('This .csv file could not be parsed. Check the delimiter and row formatting, then try again.') from error

    total_columns = max((len(row) for row in rows), default=0)
    sheet_name = Path(uploaded_file.name or '').stem or 'CSV Data'
    return [
        {
            'name': sheet_name,
            'raw_rows': rows,
            'total_columns': total_columns,
        }
    ]


def _decode_csv_content(file_bytes):
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252'):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError('This .csv file could not be decoded. Save it as UTF-8 CSV and try again.')


def _select_excel_sheet(workbook_data, sheet_name=None):
    if sheet_name:
        for sheet in workbook_data:
            if sheet['name'] == sheet_name or sheet['name'].lower() == sheet_name.lower():
                return sheet
        raise ValueError(f'The sheet or table "{sheet_name}" was not found in the uploaded file.')

    for sheet in workbook_data:
        if _sheet_has_values(sheet['raw_rows']):
            return sheet

    return workbook_data[0]


def _parse_sheet_rows(sheet_data, max_preview_columns=60, allow_empty=False):
    raw_rows = sheet_data['raw_rows']
    total_columns = sheet_data['total_columns']
    first_non_empty_index = next((index for index, row in enumerate(raw_rows) if any(str(cell).strip() for cell in row)), None)
    if first_non_empty_index is None:
        if allow_empty:
            return {
                'has_data': False,
                'headers': [],
                'cleaned_rows': [],
                'total_rows': 0,
                'total_columns': 0,
            }
        raise ValueError(f'The sheet "{sheet_data["name"]}" is empty. Add at least one header row and one data row to preview it.')

    header_row = raw_rows[first_non_empty_index]
    data_rows = raw_rows[first_non_empty_index + 1:]
    visible_columns = min(max(total_columns, len(header_row)), max_preview_columns)
    headers = _build_unique_headers(header_row[:visible_columns], visible_columns)

    cleaned_rows = []
    for row in data_rows:
        trimmed = [row[index] if index < len(row) else '' for index in range(visible_columns)]
        if any(str(cell).strip() for cell in trimmed):
            cleaned_rows.append(trimmed)

    return {
        'has_data': True,
        'headers': headers,
        'cleaned_rows': cleaned_rows,
        'total_rows': len(cleaned_rows),
        'total_columns': visible_columns,
    }


def _serialize_sheet_preview(file_name, sheet_name, parsed_sheet, max_preview_rows):
    preview_rows = parsed_sheet['cleaned_rows'][:max_preview_rows]
    return {
        'file_name': file_name,
        'sheet_name': sheet_name,
        'headers': parsed_sheet['headers'],
        'rows': preview_rows,
        'total_rows': parsed_sheet['total_rows'],
        'total_columns': parsed_sheet['total_columns'],
        'preview_truncated': parsed_sheet['total_rows'] > max_preview_rows,
        'max_preview_rows': max_preview_rows,
    }


def _summarize_excel_sheet(sheet_data, selected_sheet_name, max_preview_columns):
    parsed_sheet = _parse_sheet_rows(sheet_data, max_preview_columns=max_preview_columns, allow_empty=True)
    return {
        'name': sheet_data['name'],
        'has_data': parsed_sheet['has_data'],
        'total_rows': parsed_sheet['total_rows'],
        'total_columns': parsed_sheet['total_columns'],
        'is_selected': sheet_data['name'] == selected_sheet_name,
    }


def _build_sheet_validation_feedback(headers, cleaned_rows, mapping_profile_key=None, column_mapping=None):
    if mapping_profile_key:
        profile = EXCEL_VALIDATION_PROFILE_MAP.get(mapping_profile_key)
        if not profile:
            raise ValueError('The selected upload template is not supported.')
        _, matched_columns = _detect_excel_validation_profile(headers)
    else:
        profile, matched_columns = _detect_excel_validation_profile(headers)

    suggested_column_mapping = {}
    applied_column_mapping = {}
    sanitized_column_mapping = {}
    missing_required_columns = []
    missing_value_messages = []
    invalid_format_messages = []
    row_issue_map = {}
    cell_issue_map = {}

    if profile:
        sanitized_column_mapping = _sanitize_column_mapping(profile, headers, column_mapping)
        suggested_column_mapping = {
            canonical_name: (matched_columns.get(canonical_name) or {}).get('header', '')
            for canonical_name in profile['required_columns']
        }
        matched_columns = _merge_column_mapping_matches(headers, profile, suggested_column_mapping, sanitized_column_mapping)
        applied_column_mapping = {
            canonical_name: (matched_columns.get(canonical_name) or {}).get('header', '')
            for canonical_name in profile['required_columns']
        }
        for canonical_name in profile['required_columns']:
            if canonical_name not in matched_columns:
                missing_required_columns.append(canonical_name)

        for canonical_name, match in matched_columns.items():
            column_index = match['index']
            missing_value_count = 0
            for row_index, row in enumerate(cleaned_rows):
                cell_value = row[column_index] if column_index < len(row) else ''
                if cell_value:
                    continue
                missing_value_count += 1
                _record_sheet_cell_issue(
                    row_issue_map,
                    cell_issue_map,
                    row_index=row_index,
                    column_index=column_index,
                    tone='error',
                    message=f'Missing value for "{canonical_name}".',
                )

            if missing_value_count:
                missing_value_messages.append(
                    f'Column "{canonical_name}" has {missing_value_count} missing value{"s" if missing_value_count != 1 else ""}.'
                )

        if profile['key'] == 'visitor' and 'Date' in matched_columns:
            date_index = matched_columns['Date']['index']
            invalid_date_count = sum(
                1
                for row in cleaned_rows
                if date_index < len(row)
                and row[date_index]
                and not _is_valid_date_value(row[date_index])
            )
            if invalid_date_count:
                invalid_format_messages.append(f'Column "Date" has {invalid_date_count} invalid entr{"y" if invalid_date_count == 1 else "ies"}.')
            for row_index, row in enumerate(cleaned_rows):
                if date_index >= len(row) or not row[date_index] or _is_valid_date_value(row[date_index]):
                    continue
                _record_sheet_cell_issue(
                    row_issue_map,
                    cell_issue_map,
                    row_index=row_index,
                    column_index=date_index,
                    tone='error',
                    message='Invalid date format.',
                )

        if profile['key'] == 'record' and 'Category' in matched_columns:
            category_index = matched_columns['Category']['index']
            invalid_category_count = sum(
                1
                for row in cleaned_rows
                if category_index < len(row)
                and row[category_index]
                and row[category_index].strip().lower() not in VALID_RECORD_CATEGORIES
            )
            if invalid_category_count:
                invalid_format_messages.append(
                    f'Column "Category" has {invalid_category_count} invalid entr{"y" if invalid_category_count == 1 else "ies"}.'
                )
            for row_index, row in enumerate(cleaned_rows):
                if (
                    category_index >= len(row)
                    or not row[category_index]
                    or row[category_index].strip().lower() in VALID_RECORD_CATEGORIES
                ):
                    continue
                _record_sheet_cell_issue(
                    row_issue_map,
                    cell_issue_map,
                    row_index=row_index,
                    column_index=category_index,
                    tone='error',
                    message='Invalid category value.',
                )

    for column_index, header in enumerate(headers):
        if 'email' not in _normalize_header_key(header):
            continue
        invalid_email_count = sum(
            1
            for row in cleaned_rows
            if column_index < len(row)
            and row[column_index]
            and not EMAIL_PATTERN.match(row[column_index])
        )
        if invalid_email_count:
            invalid_format_messages.append(
                f'Column "{header}" has {invalid_email_count} invalid entr{"y" if invalid_email_count == 1 else "ies"}.'
            )
        for row_index, row in enumerate(cleaned_rows):
            if column_index >= len(row) or not row[column_index] or EMAIL_PATTERN.match(row[column_index]):
                continue
            _record_sheet_cell_issue(
                row_issue_map,
                cell_issue_map,
                row_index=row_index,
                column_index=column_index,
                tone='error',
                message=f'Invalid email in "{header}".',
            )

    duplicate_count, duplicate_basis, duplicate_row_indexes = _count_duplicate_rows(headers, cleaned_rows, profile, matched_columns)
    for row_index in duplicate_row_indexes:
        _record_sheet_row_issue(
            row_issue_map,
            row_index=row_index,
            tone='warning',
            message=f'Duplicate row based on {duplicate_basis}.',
        )

    issue_count = len(missing_required_columns) + len(missing_value_messages) + len(invalid_format_messages) + (1 if duplicate_count else 0)

    if duplicate_count:
        duplicate_message = f'{duplicate_count} duplicate row(s) detected based on {duplicate_basis}.'
    else:
        duplicate_message = 'No duplicate entries were detected in the current sheet preview.'

    if profile:
        template_message = (
            f'Validation ran against the {profile["label"]} template using your selected column mapping.'
            if mapping_profile_key
            else f'Validation ran against the {profile["label"]} template.'
        )
        required_columns = list(profile['required_columns'].keys())
    else:
        template_message = 'No known upload template was detected, so only generic duplicate and email checks were applied.'
        required_columns = []

    return {
        'profile_key': profile['key'] if profile else 'generic',
        'profile_label': profile['label'] if profile else 'Generic Sheet',
        'template_message': template_message,
        'required_columns': required_columns,
        'missing_required_columns': missing_required_columns,
        'missing_value_messages': missing_value_messages,
        'invalid_format_messages': invalid_format_messages,
        'duplicate_count': duplicate_count,
        'duplicate_row_indexes': duplicate_row_indexes,
        'duplicate_message': duplicate_message,
        'issue_count': issue_count,
        'row_issues': _serialize_row_issues(row_issue_map),
        'cell_issues': list(cell_issue_map.values()),
        'mapping_options': headers,
        'suggested_column_mapping': suggested_column_mapping,
        'applied_column_mapping': applied_column_mapping,
        'available_templates': _serialize_excel_template_options(),
    }


def _attach_submission_version_context(preview_payload, actor, original_filename):
    sheets = preview_payload.get('sheets') or []
    selected_sheet_name = preview_payload.get('selected_sheet_name') or preview_payload.get('sheet_name') or ''

    for sheet in sheets:
        version_number = _get_next_data_submission_version(actor, original_filename, sheet.get('name', ''))
        sheet['next_version_number'] = version_number
        sheet['next_version_label'] = f'v{version_number}'

    preview_payload['next_version_number'] = _get_next_data_submission_version(actor, original_filename, selected_sheet_name)
    preview_payload['next_version_label'] = f'v{preview_payload["next_version_number"]}'


def _parse_mapping_profile_key(raw_value):
    profile_key = str(raw_value or '').strip().lower()
    return profile_key if profile_key in EXCEL_VALIDATION_PROFILE_MAP else None


def _parse_submission_column_mapping(raw_value):
    if isinstance(raw_value, dict):
        payload = raw_value
    elif raw_value in {None, ''}:
        return {}
    else:
        try:
            payload = json.loads(raw_value)
        except (TypeError, ValueError) as error:
            raise ValueError('Column mapping must be valid JSON.') from error

    if not isinstance(payload, dict):
        raise ValueError('Column mapping must be a JSON object.')

    normalized = {}
    for key, value in payload.items():
        column_name = str(key or '').strip()
        mapped_header = str(value or '').strip()
        if column_name:
            normalized[column_name] = mapped_header
    return normalized


def _get_next_data_submission_version(actor, original_filename, sheet_name):
    latest_version = (
        DataSubmission.objects.filter(
            submitted_by=actor,
            original_filename=original_filename,
            sheet_name=sheet_name,
        )
        .order_by('-version_number')
        .values_list('version_number', flat=True)
        .first()
    )
    return (latest_version or 0) + 1


def _sanitize_column_mapping(profile, headers, column_mapping):
    if not profile:
        return {}

    allowed_fields = set(profile['required_columns'].keys())
    available_headers = set(headers)
    sanitized = {}
    for field_name, mapped_header in (column_mapping or {}).items():
        if field_name not in allowed_fields:
            continue
        if mapped_header and mapped_header in available_headers:
            sanitized[field_name] = mapped_header
    return sanitized


def _merge_column_mapping_matches(headers, profile, suggested_column_mapping, sanitized_column_mapping):
    matches = {}
    for canonical_name in profile['required_columns']:
        resolved_header = sanitized_column_mapping.get(canonical_name) or suggested_column_mapping.get(canonical_name) or ''
        if not resolved_header:
            continue
        try:
            column_index = headers.index(resolved_header)
        except ValueError:
            continue
        matches[canonical_name] = {
            'index': column_index,
            'header': resolved_header,
        }
    return matches


def _serialize_excel_template_options():
    return [
        {
            'key': template_key,
            'label': definition['label'],
            'required_columns': definition['headers'],
            'filename': definition['filename'],
        }
        for template_key, definition in EXCEL_TEMPLATE_DEFINITIONS.items()
    ]


def _build_submission_template_workbook(template_definition):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = template_definition['sheet_name']
    worksheet.append(template_definition['headers'])
    for row in template_definition['sample_rows']:
        worksheet.append(row)

    notes_sheet = workbook.create_sheet(title='Instructions')
    notes_sheet.append(['How to use this template'])
    for note in template_definition['notes']:
        notes_sheet.append([note])

    for column_index, header in enumerate(template_definition['headers'], start=1):
        worksheet.cell(row=1, column=column_index).font = Font(bold=True)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(18, len(header) + 4)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _detect_excel_validation_profile(headers):
    best_profile = None
    best_matches = {}
    best_score = 0

    for profile in EXCEL_VALIDATION_PROFILES:
        matches = {}
        for canonical_name, aliases in profile['required_columns'].items():
            match = _find_matching_header(headers, aliases)
            if match is not None:
                matches[canonical_name] = match
        if len(matches) > best_score:
            best_profile = profile
            best_matches = matches
            best_score = len(matches)

    if best_score < 2:
        return None, {}

    return best_profile, best_matches


def _find_matching_header(headers, aliases):
    normalized_aliases = {_normalize_header_key(alias) for alias in aliases}
    for index, header in enumerate(headers):
        if _normalize_header_key(header) in normalized_aliases:
            return {'index': index, 'header': header}
    return None


def _count_duplicate_rows(headers, cleaned_rows, profile, matched_columns):
    if not cleaned_rows:
        return 0, 'the current rows', []

    duplicate_indexes = []
    duplicate_basis = 'the full row'
    if profile:
        matched_duplicate_indexes = []
        matched_duplicate_headers = []
        for canonical_name in profile['duplicate_columns']:
            match = matched_columns.get(canonical_name)
            if match is None:
                continue
            matched_duplicate_indexes.append(match['index'])
            matched_duplicate_headers.append(canonical_name)
        if len(matched_duplicate_indexes) >= 2:
            duplicate_indexes = matched_duplicate_indexes
            duplicate_basis = ' + '.join(matched_duplicate_headers)

    signatures = {}
    duplicate_count = 0
    duplicate_row_indexes = []
    for row_index, row in enumerate(cleaned_rows):
        if duplicate_indexes:
            values = [row[index].strip().lower() if index < len(row) else '' for index in duplicate_indexes]
        else:
            values = [str(cell).strip().lower() for cell in row[:len(headers)]]
        if not any(values):
            continue
        signature = tuple(values)
        signatures.setdefault(signature, []).append(row_index)

    for row_indexes in signatures.values():
        if len(row_indexes) > 1:
            duplicate_count += len(row_indexes) - 1
            duplicate_row_indexes.extend(row_indexes)

    return duplicate_count, duplicate_basis, sorted(set(duplicate_row_indexes))


def _record_sheet_cell_issue(row_issue_map, cell_issue_map, row_index, column_index, tone, message):
    key = (row_index, column_index)
    existing = cell_issue_map.get(key)
    resolved_tone = _merge_issue_tone(existing.get('tone') if existing else None, tone)
    cell_issue_map[key] = {
        'row_index': row_index,
        'column_index': column_index,
        'tone': resolved_tone,
        'message': message,
    }
    _record_sheet_row_issue(row_issue_map, row_index=row_index, tone=tone, message=message)


def _record_sheet_row_issue(row_issue_map, row_index, tone, message):
    existing = row_issue_map.get(row_index)
    if existing is None:
        row_issue_map[row_index] = {
            'row_index': row_index,
            'tone': tone,
            'messages': [message],
        }
        return

    existing['tone'] = _merge_issue_tone(existing.get('tone'), tone)
    if message not in existing['messages']:
        existing['messages'].append(message)


def _serialize_row_issues(row_issue_map):
    return [row_issue_map[row_index] for row_index in sorted(row_issue_map)]


def _merge_issue_tone(current_tone, next_tone):
    tone_weight = {
        'neutral': 0,
        'warning': 1,
        'error': 2,
    }
    current = current_tone or 'neutral'
    upcoming = next_tone or 'neutral'
    return upcoming if tone_weight.get(upcoming, 0) >= tone_weight.get(current, 0) else current


def _normalize_submission_workspace(headers, preview_rows):
    raw_headers = headers if isinstance(headers, list) else []
    raw_rows = preview_rows if isinstance(preview_rows, list) else []
    inferred_columns = max(
        len(raw_headers),
        max((len(row) for row in raw_rows if isinstance(row, list)), default=0),
    )

    if inferred_columns <= 0:
        return [], []

    normalized_headers = _build_unique_headers(
        [_normalize_excel_value(value) for value in raw_headers[:inferred_columns]],
        inferred_columns,
    )

    normalized_rows = []
    for raw_row in raw_rows:
        if not isinstance(raw_row, list):
            continue
        normalized_row = [
            _normalize_excel_value(raw_row[index] if index < len(raw_row) else '')
            for index in range(inferred_columns)
        ]
        if any(normalized_row):
            normalized_rows.append(normalized_row)

    return normalized_headers, normalized_rows


def _build_submission_review_history_payload(submission):
    payload = []
    for raw_entry in reversed(list(submission.review_history or [])):
        entry = raw_entry if isinstance(raw_entry, dict) else {}
        payload.append(
            {
                'action_key': str(entry.get('action_key') or 'updated').strip() or 'updated',
                'title': str(entry.get('title') or 'Submission updated').strip() or 'Submission updated',
                'summary': str(entry.get('summary') or 'No summary was recorded for this action.').strip() or 'No summary was recorded for this action.',
                'details': entry.get('details') if isinstance(entry.get('details'), dict) else {},
                'changed_at': entry.get('changed_at'),
                'changed_by_id': entry.get('changed_by_id'),
                'changed_by_name': str(entry.get('changed_by_name') or '').strip(),
                'changed_by_email': str(entry.get('changed_by_email') or '').strip(),
                'changed_by_label': (
                    str(entry.get('changed_by_name') or '').strip()
                    or str(entry.get('changed_by_email') or '').strip()
                    or 'System'
                ),
            }
        )
    return payload


def _get_previous_submission_version(submission):
    return (
        DataSubmission.objects.filter(
            submitted_by_id=submission.submitted_by_id,
            original_filename=submission.original_filename,
            sheet_name=submission.sheet_name,
        )
        .exclude(pk=submission.pk)
        .filter(version_number__lt=submission.version_number)
        .order_by('-version_number', '-created_at', '-id')
        .first()
    )


def _summarize_submission_diff(previous_headers, previous_rows, next_headers, next_rows):
    previous_headers, previous_rows = _normalize_submission_workspace(previous_headers, previous_rows)
    next_headers, next_rows = _normalize_submission_workspace(next_headers, next_rows)
    items = []
    header_change_count = 0
    cell_changed_count = 0
    row_added_count = max(len(next_rows) - len(previous_rows), 0)
    row_removed_count = max(len(previous_rows) - len(next_rows), 0)
    max_items = 12

    for column_index in range(max(len(previous_headers), len(next_headers))):
        before_value = previous_headers[column_index] if column_index < len(previous_headers) else ''
        after_value = next_headers[column_index] if column_index < len(next_headers) else ''
        if before_value == after_value:
            continue
        header_change_count += 1
        if len(items) < max_items:
            items.append(
                {
                    'type': 'header',
                    'label': f'Header {column_index + 1}',
                    'before': before_value,
                    'after': after_value,
                }
            )

    shared_row_count = min(len(previous_rows), len(next_rows))
    for row_index in range(shared_row_count):
        previous_row = previous_rows[row_index]
        next_row = next_rows[row_index]
        shared_column_count = max(len(previous_row), len(next_row))
        for column_index in range(shared_column_count):
            before_value = previous_row[column_index] if column_index < len(previous_row) else ''
            after_value = next_row[column_index] if column_index < len(next_row) else ''
            if before_value == after_value:
                continue
            cell_changed_count += 1
            if len(items) < max_items:
                items.append(
                    {
                        'type': 'cell',
                        'label': f'Row {row_index + 1}, {next_headers[column_index] if column_index < len(next_headers) else f"Column {column_index + 1}"}',
                        'before': before_value,
                        'after': after_value,
                    }
                )

    for row_index in range(shared_row_count, len(next_rows)):
        if len(items) >= max_items:
            break
        items.append(
            {
                'type': 'row_added',
                'label': f'New row {row_index + 1}',
                'before': '',
                'after': ' | '.join(value for value in next_rows[row_index] if value) or 'Added row',
            }
        )

    for row_index in range(shared_row_count, len(previous_rows)):
        if len(items) >= max_items:
            break
        items.append(
            {
                'type': 'row_removed',
                'label': f'Removed row {row_index + 1}',
                'before': ' | '.join(value for value in previous_rows[row_index] if value) or 'Removed row',
                'after': '',
            }
        )

    summary_parts = []
    if header_change_count:
        summary_parts.append(f'{header_change_count} header change(s)')
    if cell_changed_count:
        summary_parts.append(f'{cell_changed_count} edited cell(s)')
    if row_added_count:
        summary_parts.append(f'{row_added_count} row(s) added')
    if row_removed_count:
        summary_parts.append(f'{row_removed_count} row(s) removed')

    return {
        'summary': ', '.join(summary_parts) if summary_parts else 'No visible row or header differences were detected.',
        'metrics': {
            'header_change_count': header_change_count,
            'cell_changed_count': cell_changed_count,
            'row_added_count': row_added_count,
            'row_removed_count': row_removed_count,
        },
        'items': items,
    }


def _build_submission_content_diff(previous_headers, previous_rows, next_headers, next_rows, max_items=12):
    diff = _summarize_submission_diff(previous_headers, previous_rows, next_headers, next_rows)
    return {
        'summary': diff['summary'],
        'metrics': diff['metrics'],
        'items': diff['items'][:max_items],
    }


def _build_submission_version_comparison(submission, headers=None, preview_rows=None):
    previous_submission = _get_previous_submission_version(submission)
    if not previous_submission:
        return {
            'has_previous_version': False,
            'current_submission_id': submission.pk,
            'current_version_label': f'v{submission.version_number}',
            'summary': 'No earlier version exists for this file yet.',
            'metrics': {
                'header_change_count': 0,
                'cell_changed_count': 0,
                'row_added_count': 0,
                'row_removed_count': 0,
            },
            'items': [],
        }

    comparison = _build_submission_content_diff(
        previous_submission.headers,
        previous_submission.preview_rows,
        submission.headers if headers is None else headers,
        submission.preview_rows if preview_rows is None else preview_rows,
    )
    return {
        'has_previous_version': True,
        'previous_submission_id': previous_submission.pk,
        'previous_version_label': f'v{previous_submission.version_number}',
        'previous_status': previous_submission.status,
        'previous_created_at': previous_submission.created_at,
        'current_submission_id': submission.pk,
        'current_version_label': f'v{submission.version_number}',
        'current_status': submission.status,
        'current_created_at': submission.created_at,
        'summary': comparison['summary'],
        'metrics': comparison['metrics'],
        'items': comparison['items'],
    }


def _build_submission_workspace_update_summary(workspace_diff, *, mapping_profile_changed=False, mapping_change_count=0):
    summary_parts = []
    if workspace_diff['summary'] and workspace_diff['summary'] != 'No visible row or header differences were detected.':
        summary_parts.append(workspace_diff['summary'])
    if mapping_profile_changed:
        summary_parts.append('template mapping changed')
    if mapping_change_count:
        summary_parts.append(f'{mapping_change_count} mapped field change(s)')
    if not summary_parts:
        return 'Saved the preview workspace without changing any visible rows or mappings.'
    return f'Saved preview fixes with {", ".join(summary_parts)}.'


def _resolve_submission_workspace_payload(submission, headers=None, preview_rows=None, mapping_profile_key=None, column_mapping=None):
    resolved_headers, resolved_rows = _normalize_submission_workspace(
        submission.headers if headers is None else headers,
        submission.preview_rows if preview_rows is None else preview_rows,
    )
    validation = _build_sheet_validation_feedback(
        resolved_headers,
        resolved_rows,
        mapping_profile_key=mapping_profile_key if mapping_profile_key is not None else submission.mapping_profile_key,
        column_mapping=column_mapping if column_mapping is not None else submission.column_mapping,
    )
    resolved_profile_key = validation.get('profile_key') if validation.get('profile_key') != 'generic' else ''
    resolved_column_mapping = validation.get('applied_column_mapping', {})

    return {
        'file_name': submission.original_filename,
        'sheet_name': submission.sheet_name or 'Submitted worksheet',
        'headers': resolved_headers,
        'rows': resolved_rows,
        'total_rows': len(resolved_rows),
        'total_columns': len(resolved_headers),
        'preview_truncated': False,
        'max_preview_rows': len(resolved_rows),
        'mapping_profile_key': resolved_profile_key,
        'column_mapping': resolved_column_mapping,
        'validation': validation,
    }


def _build_submission_workspace_response(submission, request, workspace_preview=None):
    preview_payload = workspace_preview or _resolve_submission_workspace_payload(submission)
    priority_metadata = build_submission_priority_metadata(submission)
    serializer = DataSubmissionSerializer(submission, context={'request': request})
    return {
        'submission': serializer.data,
        'workspace_preview': {
            **preview_payload,
            'source_submission_id': submission.id,
            'source_submission_status': submission.status,
            'source_submission_sender': submission.submitted_by.name or submission.submitted_by.email or submission.submitted_by.username,
            'source_submission_feedback': submission.admin_feedback or '',
            'source_submission_created_at': submission.created_at,
            'source_submission_updated_at': submission.updated_at,
            'priority': priority_metadata,
            'review_history': _build_submission_review_history_payload(submission),
            'version_comparison': _build_submission_version_comparison(
                submission,
                headers=preview_payload.get('headers'),
                preview_rows=preview_payload.get('rows'),
            ),
        },
    }


def _apply_submission_status_update(submission, next_status, admin_feedback, actor=None, request=None, log_action_name='Updated Excel submission status'):
    allowed_statuses = {
        DataSubmission.STATUS_PENDING,
        DataSubmission.STATUS_IN_REVIEW,
        DataSubmission.STATUS_NEEDS_REVISION,
        DataSubmission.STATUS_APPROVED,
        DataSubmission.STATUS_REJECTED,
    }
    if next_status not in allowed_statuses:
        raise ValueError('Invalid submission status.')

    if next_status in {DataSubmission.STATUS_NEEDS_REVISION, DataSubmission.STATUS_REJECTED} and not admin_feedback:
        raise ValueError('Admin feedback is required when a submission needs revision or is rejected.')

    previous_status = submission.status
    previous_feedback = submission.admin_feedback or ''
    status_changed = previous_status != next_status
    feedback_changed = previous_feedback != admin_feedback

    if not status_changed and not feedback_changed:
        return False

    submission.status = next_status
    submission.admin_feedback = admin_feedback
    append_submission_review_event(
        submission,
        actor=actor,
        action_key='status_updated',
        title='Review decision updated',
        summary=_build_submission_status_update_summary(
            previous_status,
            next_status,
            feedback_changed=feedback_changed,
        ),
        details={
            'previous_status': previous_status,
            'next_status': next_status,
            'previous_feedback': previous_feedback,
            'next_feedback': admin_feedback,
        },
        mark_reviewed=actor is not None and getattr(actor, 'is_authenticated', False),
    )
    update_fields = ['status', 'admin_feedback', 'review_history', 'updated_at']
    if actor is not None and getattr(actor, 'is_authenticated', False):
        update_fields.extend(['last_reviewed_by', 'last_reviewed_at'])
    submission.save(update_fields=update_fields)
    if actor is not None:
        create_submission_notification(submission, actor=actor)
        log_action(
            actor,
            log_action_name,
            (
                f'Updated data file "{submission.original_filename}" from {previous_status} to {next_status}. '
                f'Feedback: {admin_feedback}'
            ).strip(),
            request,
        )
    return True


def _build_submission_status_update_summary(previous_status, next_status, *, feedback_changed=False):
    if previous_status != next_status and feedback_changed:
        return f'Changed the review status from {previous_status} to {next_status} and updated the admin feedback.'
    if previous_status != next_status:
        return f'Changed the review status from {previous_status} to {next_status}.'
    if feedback_changed:
        return f'Updated the admin feedback while keeping the file in {next_status}.'
    return 'Updated the review details.'


def _parse_submission_workspace_request_payload(raw_payload):
    if not isinstance(raw_payload, dict):
        raise ValueError('Invalid workspace payload.')

    headers = raw_payload.get('headers')
    preview_rows = raw_payload.get('preview_rows', raw_payload.get('rows'))

    if not isinstance(headers, list):
        raise ValueError('Headers must be provided as an array.')
    if not isinstance(preview_rows, list):
        raise ValueError('Preview rows must be provided as an array of rows.')
    if len(headers) > 60:
        raise ValueError('A submission preview can contain at most 60 columns.')
    if len(preview_rows) > 1000:
        raise ValueError('A submission preview can contain at most 1000 rows.')

    for row in preview_rows:
        if not isinstance(row, list):
            raise ValueError('Each preview row must be an array of cell values.')

    mapping_profile_key = _parse_mapping_profile_key(raw_payload.get('mapping_profile_key'))
    column_mapping = _parse_submission_column_mapping(raw_payload.get('column_mapping'))
    return {
        'headers': headers,
        'preview_rows': preview_rows,
        'mapping_profile_key': mapping_profile_key,
        'column_mapping': column_mapping,
    }


def _parse_submission_id_list(raw_value):
    if isinstance(raw_value, (list, tuple)):
        candidate_values = list(raw_value)
    elif isinstance(raw_value, str):
        candidate_values = [value.strip() for value in raw_value.split(',') if value.strip()]
    else:
        raise ValueError('Select at least one submission first.')

    parsed_ids = []
    for value in candidate_values:
        try:
            submission_id = int(value)
        except (TypeError, ValueError) as error:
            raise ValueError('Submission ids must be valid integers.') from error
        if submission_id > 0 and submission_id not in parsed_ids:
            parsed_ids.append(submission_id)

    if not parsed_ids:
        raise ValueError('Select at least one submission first.')
    return parsed_ids


def _parse_user_id_list(raw_value):
    if isinstance(raw_value, (list, tuple)):
        candidate_values = list(raw_value)
    elif isinstance(raw_value, str):
        candidate_values = [value.strip() for value in raw_value.split(',') if value.strip()]
    else:
        raise ValueError('Choose at least one user to notify.')

    parsed_ids = []
    for value in candidate_values:
        try:
            user_id = int(value)
        except (TypeError, ValueError) as error:
            raise ValueError('User ids must be valid integers.') from error
        if user_id > 0 and user_id not in parsed_ids:
            parsed_ids.append(user_id)

    if not parsed_ids:
        raise ValueError('Choose at least one user to notify.')
    return parsed_ids


def _parse_establishment_name_list(raw_value):
    if isinstance(raw_value, (list, tuple)):
        candidate_values = list(raw_value)
    elif isinstance(raw_value, str):
        candidate_values = [value.strip() for value in raw_value.split(',') if value.strip()]
    else:
        raise ValueError('Choose at least one establishment to notify.')

    parsed_names = []
    for value in candidate_values:
        establishment_name = str(value or '').strip()
        if establishment_name and establishment_name not in parsed_names:
            parsed_names.append(establishment_name)

    if not parsed_names:
        raise ValueError('Choose at least one establishment to notify.')
    return parsed_names


def _parse_data_request_due_date(raw_value):
    raw_text = str(raw_value or '').strip()
    if not raw_text:
        return None
    parsed_due_date = parse_date(raw_text)
    if not parsed_due_date:
        raise ValueError('Due date must be a valid date.')
    if parsed_due_date < timezone.localdate():
        raise ValueError('Due date cannot be in the past.')
    return parsed_due_date


def _get_data_request_target_queryset():
    return User.objects.filter(
        status=UserStatusChoices.ACTIVE,
        deleted_at__isnull=True,
    ).order_by('office', 'name', 'email', 'id')


def _serialize_data_request_target(user):
    return {
        'id': user.pk,
        'name': user.name,
        'username': user.username,
        'email': user.email,
        'office': user.office,
        'role': user.role,
    }


def _resolve_data_request_recipients(request):
    target_users = [
        user
        for user in _get_data_request_target_queryset()
        if user.pk != request.user.pk and not can_access_admin_dashboard(user)
    ]
    target_by_id = {user.pk: user for user in target_users}
    recipient_scope = str(request.data.get('recipient_scope', '')).strip().lower()

    if recipient_scope == 'all':
        return target_users

    if recipient_scope in {'establishment', 'establishments'}:
        requested_establishments = _parse_establishment_name_list(
            request.data.get('establishments') or request.data.get('establishment_names')
        )
        users_by_establishment = {}
        for user in target_users:
            establishment_name = str(user.office or '').strip()
            if not establishment_name:
                continue
            users_by_establishment.setdefault(establishment_name, []).append(user)

        missing_establishments = [
            establishment_name
            for establishment_name in requested_establishments
            if establishment_name not in users_by_establishment
        ]
        if missing_establishments:
            raise ValueError('One or more selected establishments cannot receive data requests.')

        recipients = []
        seen_user_ids = set()
        for establishment_name in requested_establishments:
            for user in users_by_establishment[establishment_name]:
                if user.pk not in seen_user_ids:
                    recipients.append(user)
                    seen_user_ids.add(user.pk)
        return recipients

    requested_ids = _parse_user_id_list(request.data.get('recipient_ids'))
    missing_ids = [user_id for user_id in requested_ids if user_id not in target_by_id]
    if missing_ids:
        raise ValueError('One or more selected users cannot receive data requests.')
    return [target_by_id[user_id] for user_id in requested_ids]


def _build_unique_export_filename(used_names, candidate_name):
    base_name = Path(candidate_name).stem or 'submission'
    suffix = Path(candidate_name).suffix or '.xlsx'
    next_name = f'{base_name}{suffix}'
    index = 2
    while next_name in used_names:
        next_name = f'{base_name}-{index}{suffix}'
        index += 1
    used_names.add(next_name)
    return next_name


def _sheet_has_values(raw_rows):
    for row in raw_rows:
        if any(str(cell).strip() for cell in row):
            return True
    return False


def _normalize_header_key(value):
    return HEADER_PATTERN.sub(' ', str(value or '').strip().lower()).strip()


def _is_valid_date_value(value):
    text_value = str(value or '').strip()
    if not text_value:
        return False
    if parse_datetime(text_value) or parse_date(text_value):
        return True
    for date_format in (
        '%m/%d/%Y',
        '%m/%d/%y',
        '%Y-%m-%d',
        '%Y/%m/%d',
        '%b %d %Y',
        '%b %d, %Y',
        '%B %d %Y',
        '%B %d, %Y',
        '%d %b %Y',
        '%d %B %Y',
    ):
        try:
            return bool(datetime.strptime(text_value, date_format))
        except ValueError:
            continue
    return False


def _build_unique_headers(header_row, total_columns):
    seen = {}
    headers = []
    for index in range(total_columns):
        raw_value = header_row[index] if index < len(header_row) else ''
        base = str(raw_value).strip() or f'Column {index + 1}'
        count = seen.get(base, 0) + 1
        seen[base] = count
        headers.append(base if count == 1 else f'{base} ({count})')
    return headers


def _normalize_excel_value(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
