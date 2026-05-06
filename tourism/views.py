from calendar import month_name
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Q
from rest_framework import status, viewsets
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from accounts.permissions import CanDeleteDataPermission, CanEditDataPermission, IsActiveSystemUser
from audit.services import log_action, soft_delete_to_recycle
from citosis_pro.common import RecycleItemTypeChoices
from tourism.models import TourismRecord, Visitor
from tourism.serializers import TourismRecordSerializer, VisitorSerializer


class TourismRecordViewSet(viewsets.ModelViewSet):
    serializer_class = TourismRecordSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    queryset = TourismRecord.objects.select_related('created_by', 'updated_by').all()

    def get_permissions(self):
        if self.action == 'destroy':
            permission_classes = [CanDeleteDataPermission]
        elif self.action in {'create', 'update', 'partial_update'}:
            permission_classes = [CanEditDataPermission]
        else:
            permission_classes = [IsActiveSystemUser]
        return [permission() for permission in permission_classes]

    def get_queryset(self):
        queryset = TourismRecord.objects.select_related('created_by', 'updated_by').order_by('-updated_at')
        search = self.request.query_params.get('search', '').strip()
        category = self.request.query_params.get('category', '').strip()
        if search:
            queryset = queryset.filter(name__icontains=search) | queryset.filter(location__icontains=search) | queryset.filter(description__icontains=search)
        if category:
            queryset = queryset.filter(category=category)
        return queryset.distinct()

    def perform_create(self, serializer):
        record = serializer.save(created_by=self.request.user)
        log_action(self.request.user, 'Created tourism record', f'Added record #{record.pk} ({record.name}).', self.request)

    def perform_update(self, serializer):
        record = serializer.save(updated_by=self.request.user)
        log_action(self.request.user, 'Updated tourism record', f'Updated record #{record.pk} ({record.name}).', self.request)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        soft_delete_to_recycle(instance, RecycleItemTypeChoices.RECORD, request.user, request=request)
        return Response(status=status.HTTP_204_NO_CONTENT)


class VisitorViewSet(viewsets.ModelViewSet):
    serializer_class = VisitorSerializer
    queryset = Visitor.objects.select_related('tourism_record', 'created_by', 'updated_by').all()

    def get_permissions(self):
        if self.action == 'destroy':
            permission_classes = [CanDeleteDataPermission]
        elif self.action in {'create', 'update', 'partial_update'}:
            permission_classes = [CanEditDataPermission]
        else:
            permission_classes = [IsActiveSystemUser]
        return [permission() for permission in permission_classes]

    def get_queryset(self):
        queryset = Visitor.objects.select_related('tourism_record', 'created_by', 'updated_by').order_by('-visit_date')
        search = self.request.query_params.get('search', '').strip()
        status_value = self.request.query_params.get('status', '').strip()
        if search:
            queryset = queryset.filter(name__icontains=search) | queryset.filter(place__icontains=search) | queryset.filter(origin__icontains=search)
        if status_value:
            queryset = queryset.filter(status=status_value)
        return queryset.distinct()

    def perform_create(self, serializer):
        visitor = serializer.save(created_by=self.request.user)
        log_action(self.request.user, 'Created visitor entry', f'Added visitor #{visitor.pk} ({visitor.name}).', self.request)

    def perform_update(self, serializer):
        visitor = serializer.save(updated_by=self.request.user)
        log_action(self.request.user, 'Updated visitor entry', f'Updated visitor #{visitor.pk} ({visitor.name}).', self.request)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        soft_delete_to_recycle(instance, RecycleItemTypeChoices.VISITOR, request.user, request=request)
        return Response(status=status.HTTP_204_NO_CONTENT)


class TourismAttractionRecordExportView(APIView):
    permission_classes = [IsActiveSystemUser]

    def get(self, request):
        month = self._parse_int(request.query_params.get('month'), default=timezone.localdate().month, minimum=1, maximum=12)
        year = self._parse_int(request.query_params.get('year'), default=timezone.localdate().year, minimum=2000, maximum=9999)

        workbook = build_tourism_attraction_record_workbook(month=month, year=year)
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = f'tourism-attraction-record-{year}-{month:02d}.xlsx'
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        log_action(
            request.user,
            'Exported tourism attraction record template',
            f'Generated tourism attraction record Excel template for {month_name[month]} {year}.',
            request,
        )
        return response

    @staticmethod
    def _parse_int(raw_value, default, minimum=None, maximum=None):
        try:
            value = int(raw_value)
        except (TypeError, ValueError):
            return default
        if minimum is not None and value < minimum:
            return default
        if maximum is not None and value > maximum:
            return default
        return value


def build_tourism_attraction_record_workbook(month, year):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Tourism Record'

    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)

    column_widths = {
        'A': 8,
        'B': 28,
        'C': 10,
        'D': 9,
        'E': 9,
        'F': 9,
        'G': 9,
        'H': 9,
        'I': 9,
        'J': 9,
        'K': 9,
        'L': 9,
        'M': 10,
        'N': 10,
        'O': 10,
        'P': 10,
        'Q': 10,
        'R': 10,
    }
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.page_setup.orientation = 'landscape'
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LEGAL
    worksheet.freeze_panes = 'A8'

    worksheet.merge_cells('A1:R1')
    worksheet['A1'] = 'Tourism Attraction Record'
    worksheet['A1'].font = title_font
    worksheet['A1'].alignment = center

    worksheet.merge_cells('A2:R2')
    worksheet['A2'] = '(This Recording Form can be used instead of just counting visitors)'
    worksheet['A2'].alignment = center

    worksheet.merge_cells('A4:J4')
    worksheet['A4'] = 'Name of Municipality: Malaybalay City, Bukidnon'
    worksheet['A4'].font = bold
    worksheet['A4'].alignment = left

    worksheet.merge_cells('K4:R4')
    worksheet['K4'] = f'Month/Year: {month_name[month]} {year}'
    worksheet['K4'].font = bold
    worksheet['K4'].alignment = left

    worksheet.merge_cells('D5:L5')
    worksheet['D5'] = 'Place of Residence Philippines'
    worksheet['D5'].font = bold
    worksheet['D5'].alignment = center

    worksheet.merge_cells('A6:A7')
    worksheet['A6'] = 'Date'
    worksheet.merge_cells('B6:B7')
    worksheet['B6'] = 'Name'
    worksheet.merge_cells('C6:C7')
    worksheet['C6'] = 'Code'

    worksheet.merge_cells('D6:F6')
    worksheet['D6'] = 'This Municipality'
    worksheet.merge_cells('G6:I6')
    worksheet['G6'] = 'Other Municipality'
    worksheet.merge_cells('J6:L6')
    worksheet['J6'] = 'Other Province'
    worksheet.merge_cells('M6:O6')
    worksheet['M6'] = 'Foreign Country Residence'
    worksheet.merge_cells('P6:R6')
    worksheet['P6'] = 'Grand Total Number of Visitors'

    subheaders = ['Male', 'Female', 'Total']
    for start_column in ('D', 'G', 'J', 'M', 'P'):
        start_index = ord(start_column) - ord('A') + 1
        for offset, label in enumerate(subheaders):
            cell = worksheet.cell(row=7, column=start_index + offset)
            cell.value = label
            cell.font = bold
            cell.alignment = center
            cell.border = border

    for row in range(6, 8):
        worksheet.row_dimensions[row].height = 24

    for cell_name in ('A6', 'B6', 'C6', 'D6', 'G6', 'J6', 'M6', 'P6'):
        worksheet[cell_name].font = bold
        worksheet[cell_name].alignment = center

    start_data_row = 8
    end_data_row = 38
    for day in range(1, 32):
        row_index = start_data_row + day - 1
        worksheet.cell(row=row_index, column=1, value=day)
        worksheet.cell(row=row_index, column=16, value=f'=SUM(D{row_index},G{row_index},J{row_index},M{row_index})')
        worksheet.cell(row=row_index, column=17, value=f'=SUM(E{row_index},H{row_index},K{row_index},N{row_index})')
        worksheet.cell(row=row_index, column=18, value=f'=SUM(F{row_index},I{row_index},L{row_index},O{row_index})')

    total_row = end_data_row + 1
    worksheet.merge_cells(f'A{total_row}:C{total_row}')
    worksheet[f'A{total_row}'] = 'TOTAL of this Month'
    worksheet[f'A{total_row}'].font = bold
    worksheet[f'A{total_row}'].alignment = center

    for column in range(4, 19):
        column_letter = get_column_letter(column)
        worksheet.cell(row=total_row, column=column, value=f'=SUM({column_letter}{start_data_row}:{column_letter}{end_data_row})')

    for row in range(6, total_row + 1):
        for column in range(1, 19):
            cell = worksheet.cell(row=row, column=column)
            cell.border = border
            if row >= start_data_row:
                cell.alignment = center if column != 2 else left

    return workbook
